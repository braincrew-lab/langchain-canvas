"""Workbook charts paired with their anchors through a marked chart gallery.

The read-only grid shows each chart as a picture at its anchor cell. The host
renders a print copy with one page per chart, each page marked
``lcx<salt>chart<k>``; the crop of that page belongs to the k-th chart box, so
no page map is predicted and no "is this a chart" guess is made.
"""

from __future__ import annotations

import hashlib
import io
import re
import shutil
import struct
import subprocess
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from PIL import Image
from test_file_preview import _chart_page, _vector_pdf, _XlsxPages

from langchain_canvas.converters import PdfSourceConverter
from langchain_canvas.replay import _PREVIEW_CACHE, _file_preview_data
from langchain_canvas.store import InMemoryCanvasStore
from langchain_canvas.xlsx_gallery import MAX_GALLERY_CHARTS, gallery_salt, xlsx_chart_gallery
from langchain_canvas.xlsx_import import xlsx_sheets_and_charts

_FIXTURES = Path(__file__).parent / "fixtures"
_EMU_PER_PX = 9525
_SHEET_MARKER = re.compile(r"^lcx([0-9a-f]{12})chart(\d+)$")


@pytest.fixture(autouse=True)
def _clear_preview_cache():
    _PREVIEW_CACHE.clear()
    yield
    _PREVIEW_CACHE.clear()


# --- builders ---------------------------------------------------------------------


def _bar_chart(ws, title: str | None = None) -> BarChart:
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    if title:
        chart.title = title
    return chart


def _fill(ws) -> None:
    ws.append(["구분", "수"])
    ws.append(["운영", 3])
    ws.append(["개발", 1])


def _save(wb: openpyxl.Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _chart_workbook() -> bytes:
    """One sheet, one chart at A15 (top 280px in the grid)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "피벗"
    _fill(ws)
    ws.add_chart(_bar_chart(ws), "A15")
    return _save(wb)


def _two_chart_workbook() -> bytes:
    """Two sheets, one titled chart each, both at A15."""
    wb = openpyxl.Workbook()
    first = wb.active
    first.title = "첫째"
    second = wb.create_sheet("둘째")
    for ws in (first, second):
        _fill(ws)
        ws.add_chart(_bar_chart(ws, title=f"차트 {ws.title}"), "A15")
    return _save(wb)


def _png(width: int = 2, height: int = 2, *, pad: int = 0) -> bytes:
    """A real PNG of ``width`` x ``height`` black pixels; ``pad`` bytes of text ride in a chunk."""

    def chunk(kind: bytes, body: bytes) -> bytes:
        crc = zlib.crc32(kind + body)
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", crc)

    rows = b"".join(b"\x00" + b"\x00\x00\x00" * width for _ in range(height))
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    padding = chunk(b"tEXt", b"pad\x00" + b"x" * pad) if pad else b""
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + padding
        + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")
    )


def _chart_title(chart) -> str:
    if chart.title is None:
        return ""
    rich = chart.title.tx.rich
    return "".join(run.t for paragraph in rich.p for run in (paragraph.r or []))


# --- a host that renders the gallery copy ------------------------------------------


class _GalleryHost(_XlsxPages):
    """A host's office renderer: one crop per marked visible sheet of the copy it is given.

    It reads the markers off the copy's sheet titles, so what it returns is
    exactly as good as the copy the replay asked for. ``rendered`` keeps every
    byte string it was asked to render.
    """

    def __init__(self, *, salt: str | None = None, png: bytes | None = None) -> None:
        super().__init__()
        self.salt = salt
        self.png = png
        self.rendered: list[bytes] = []
        self.gallery_calls = 0
        self.titles: dict[int, str] = {}
        self.charts_per_sheet: dict[int, int] = {}

    def render_pages(self, data: bytes, *, path: str, pages: list[int]):
        self.rendered.append(data)
        return super().render_pages(data, path=path, pages=pages)

    def chart_pages(self, data: bytes, *, path: str) -> list[int]:
        self.rendered.append(data)
        return super().chart_pages(data, path=path)

    def chart_images(self, data: bytes, *, path: str) -> list[dict]:
        self.rendered.append(data)
        self.gallery_calls += 1
        book = openpyxl.load_workbook(io.BytesIO(data))
        crops = []
        for sheet in book.worksheets:
            found = _SHEET_MARKER.match(sheet.title)
            if found is None or sheet.sheet_state != "visible":
                continue
            number = int(found[2])
            self.charts_per_sheet[number] = len(sheet._charts)
            self.titles[number] = _chart_title(sheet._charts[0]) if sheet._charts else ""
            crops.append({
                "chart": number,
                "salt": self.salt or found[1],
                "page": number + 1,
                "box": [50, 100, 300, 200],
                "png": self.png if self.png is not None else _png(),
            })
        return crops

    def distinct_rendered(self) -> int:
        return len({hashlib.sha256(data).hexdigest() for data in self.rendered})


def _preview(workbook: bytes, host: _GalleryHost) -> dict:
    store = InMemoryCanvasStore()
    commit = store.write_bytes("c1", "report.xlsx", workbook, "chart")
    data = _file_preview_data(store, "c1", "report.xlsx", commit.revision, [host])
    assert data is not None
    return data


def _box(image: dict) -> tuple[int, int, int, int]:
    return image["left"], image["top"], image["width"], image["height"]


# --- the gallery copy ---------------------------------------------------------------


def test_gallery_copy_gives_each_chart_its_own_marked_visible_sheet() -> None:
    data = _two_chart_workbook()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)
    salt = gallery_salt(data)

    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, entries)))

    states = {sheet.title: sheet.sheet_state for sheet in book.worksheets}
    markers = [f"lcx{salt}chart0", f"lcx{salt}chart1"]
    assert states == {
        "첫째": "hidden", "둘째": "hidden", markers[0]: "visible", markers[1]: "visible"
    }
    assert salt == hashlib.sha256(data).hexdigest()[:12]
    for marker in markers:
        sheet = book[marker]
        assert sheet["A1"].value == marker
        assert len(sheet._charts) == 1
    assert [len(book[name]._charts) for name in ("첫째", "둘째")] == [0, 0]
    assert book.active.title == markers[0]


def test_gallery_copy_is_deterministic(monkeypatch) -> None:
    import datetime

    import openpyxl.writer.excel as excel

    ticks = iter(datetime.datetime(2026, 10, 20, 9, 0, second) for second in range(10))
    clock = SimpleNamespace(
        datetime=SimpleNamespace(now=lambda tz=None: next(ticks)), timezone=datetime.timezone
    )
    monkeypatch.setattr(excel, "datetime", clock)

    data = _two_chart_workbook()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)
    first = xlsx_chart_gallery(data, entries)
    second = xlsx_chart_gallery(data, entries)

    assert first == second
    cores = [
        zipfile.ZipFile(io.BytesIO(copy)).read("docProps/core.xml") for copy in (first, second)
    ]
    assert cores[0] == cores[1]
    assert b"2026" not in cores[0]


def test_gallery_clears_print_areas_and_titles_of_hidden_sheets() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "인쇄"
    _fill(ws)
    ws.print_area = "A1:B2"
    ws.print_title_rows = "1:1"
    ws.print_title_cols = "A:A"
    ws.add_chart(_bar_chart(ws), "A15")
    data = _save(wb)
    _, entries = xlsx_sheets_and_charts(data, read_only=True)

    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, entries)))

    original = book["인쇄"]
    assert not original.print_area
    assert original.print_title_rows is None and original.print_title_cols is None


def test_gallery_refuses_entries_it_cannot_serve() -> None:
    data = _chart_workbook()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)
    box = entries[0]

    for bad in (
        [],
        [box] * (MAX_GALLERY_CHARTS + 1),
        [{**box, "sheet": 7}],
        [{**box, "sheet": None}],
        [{**box, "slot": 3}],
        [{**box, "kind": "unboxed"}],  # nothing to draw
    ):
        with pytest.raises(ValueError):
            xlsx_chart_gallery(data, bad)


def test_gallery_refuses_an_archive_over_the_safety_limits() -> None:
    from langchain_canvas.converters import UnsafeArchiveError
    data = _chart_workbook()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/blob.bin", b"\0" * 300 * 1024 * 1024)

    with pytest.raises(UnsafeArchiveError):
        xlsx_chart_gallery(bomb.getvalue(), entries)


def test_gallery_chart_size_is_clamped() -> None:
    data = _chart_workbook()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)

    def drawn(width: int, height: int) -> tuple[float, float]:
        sized = [{**entries[0], "width": width, "height": height}]
        book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, sized)))
        ext = book.active._charts[0].anchor.ext
        return ext.width / _EMU_PER_PX, ext.height / _EMU_PER_PX

    width, height = drawn(16, 19)  # a chart folded to a sliver
    assert width >= 120 and height >= 80
    assert abs(width / height - 16 / 19) < 0.05
    width, height = drawn(5000, 4000)  # a chart the size of a street
    assert width <= 1600 and height <= 1200
    assert abs(width / height - 5000 / 4000) < 0.05
    assert drawn(567, 283) == (567, 283)  # a normal chart keeps its size
    width, height = drawn(100, 20_971_500)  # a box one row-count tall and one column wide
    assert width >= 120 and height == 1200


# --- the marker crop ----------------------------------------------------------------


def test_chart_images_crop_by_marker_not_by_colored_marks() -> None:
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart1) Tj ET"
    outline = "0 0 0 RG 1 w 100 300 200 150 re S"  # no fill, no colored mark
    pages = [_chart_page(50, 300), marker + " " + outline]

    crops = PdfSourceConverter().chart_images(_vector_pdf(pages), path="gallery.pdf")

    assert [(c["chart"], c["salt"], c["page"]) for c in crops] == [(1, "0123456789ab", 2)]
    x, y, width, height = crops[0]["box"]
    # a stroke's bounds reach half a line width past the rectangle; y is from the top
    assert abs(x - 100) <= 2 and abs(y - 392) <= 2
    assert abs(width - 200) <= 2 and abs(height - 150) <= 2
    assert crops[0]["png"].startswith(b"\x89PNG")


def test_chart_images_without_markers_returns_nothing() -> None:
    pdf = _vector_pdf([_chart_page(50, 300)])  # a colored chart, no marker
    assert PdfSourceConverter().chart_images(pdf, path="gallery.pdf") == []


def test_chart_title_containing_lcx_stays_in_the_crop_box() -> None:
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart0) Tj ET"
    title = "BT /F1 20 Tf 120 470 Td (lcx Sales) Tj ET"
    outline = "0 0 0 RG 1 w 100 300 200 150 re S"

    # the title is drawn first, so its "lcx" is the first token the page text offers
    crops = PdfSourceConverter().chart_images(
        _vector_pdf([title + " " + marker + " " + outline]), path="gallery.pdf"
    )

    assert [(c["chart"], c["salt"]) for c in crops] == [(0, "0123456789ab")]
    top = crops[0]["box"][1]
    # the title sits above the outline (outline top is 392 from the page top);
    # the marker at the page top (42) is left out of the box
    assert 300 < top <= 362


def test_a_marked_page_with_nothing_else_drawn_has_no_crop() -> None:
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart0) Tj ET"
    assert PdfSourceConverter().chart_images(_vector_pdf([marker]), path="gallery.pdf") == []


def test_an_object_without_bounds_is_left_out_of_the_crop_box(monkeypatch) -> None:
    import pypdfium2.raw as pdfium_c

    real = pdfium_c.FPDFPageObj_GetBounds
    calls = []

    def _first_has_no_bounds(raw, left, bottom, right, top):
        calls.append(raw)
        return 0 if len(calls) == 1 else real(raw, left, bottom, right, top)

    monkeypatch.setattr(pdfium_c, "FPDFPageObj_GetBounds", _first_has_no_bounds)
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart0) Tj ET"
    shapes = "0 0 0 RG 1 w 300 300 200 150 re S 50 50 80 60 re S"

    crops = PdfSourceConverter().chart_images(
        _vector_pdf([marker + " " + shapes]), path="gallery.pdf"
    )

    assert len(calls) == 2  # the marker is skipped before its bounds are asked for
    x, _, width, height = crops[0]["box"]
    # only the second rectangle (50, 50, 80x60) forms the box
    assert abs(x - 50) <= 2 and abs(width - 80) <= 2 and abs(height - 60) <= 2


def test_first_page_per_marker_wins() -> None:
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart0) Tj ET"
    pages = [
        marker + " 0 0 0 RG 1 w 100 300 200 150 re S",
        marker + " 0 0 0 RG 1 w 50 50 80 60 re S",
    ]

    crops = PdfSourceConverter().chart_images(_vector_pdf(pages), path="gallery.pdf")

    assert [c["page"] for c in crops] == [1]


# --- the replay pairs crops with anchors ---------------------------------------------


def test_workbook_charts_land_at_their_anchor_via_markers() -> None:
    host = _GalleryHost()
    data = _preview(_chart_workbook(), host)

    images = data["workbook"]["sheet"][0]["images"]
    assert [i["id"] for i in images] == ["chart_0"]
    assert images[0]["src"].startswith("data:image/png;base64,")
    assert _box(images[0]) == (0, 280, 567, 283)
    assert data["chartPages"] is None
    assert "_charts" not in data["workbook"]


def test_two_sheets_each_get_their_chart_at_the_anchor() -> None:
    host = _GalleryHost()
    data = _preview(_two_chart_workbook(), host)

    first, second = data["workbook"]["sheet"]
    assert [i["id"] for i in first["images"]] == ["chart_0"]
    assert [i["id"] for i in second["images"]] == ["chart_1"]
    assert _box(first["images"][0]) == _box(second["images"][0]) == (0, 280, 567, 283)
    assert host.titles == {0: "차트 첫째", 1: "차트 둘째"}
    assert data["chartPages"] is None


def test_a_sheet_without_charts_in_front_does_not_shift_the_pairing() -> None:
    wb = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    wb.create_sheet("메모", 0)  # empty, visible, first
    host = _GalleryHost()
    data = _preview(_save(wb), host)

    memo, pivot = data["workbook"]["sheet"]
    assert not memo.get("images")
    assert [i["id"] for i in pivot["images"]] == ["chart_0"]
    assert _box(pivot["images"][0]) == (0, 280, 567, 283)
    assert data["chartPages"] is None


def test_two_charts_on_one_sheet_each_land_at_their_anchor() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    _fill(ws)
    lower = _bar_chart(ws, "아래")
    ws.add_chart(lower, "A30")  # first in the sheet, second by position
    upper = _bar_chart(ws, "위")
    ws.add_chart(upper, "A15")
    host = _GalleryHost()
    data = _preview(_save(wb), host)

    images = data["workbook"]["sheet"][0]["images"]
    assert [(i["id"], i["top"]) for i in images] == [("chart_0", 280), ("chart_1", 580)]
    assert host.titles == {0: "위", 1: "아래"}
    assert host.charts_per_sheet == {0: 1, 1: 1}
    assert data["chartPages"] is None


def test_one_preview_renders_original_and_one_gallery_copy() -> None:
    host = _GalleryHost()
    _preview(_two_chart_workbook(), host)

    assert host.distinct_rendered() == 2
    assert host.gallery_calls == 1


def test_gallery_failure_falls_back_to_chart_pages(monkeypatch) -> None:
    import langchain_canvas.xlsx_gallery as xlsx_gallery

    def _boom(data: bytes, entries: list[dict]) -> bytes:
        raise ValueError("cannot build the gallery")

    monkeypatch.setattr(xlsx_gallery, "xlsx_chart_gallery", _boom)
    host = _GalleryHost()
    data = _preview(_two_chart_workbook(), host)

    assert data["chartPages"] == [2]
    assert not any(sheet.get("images") for sheet in data["workbook"]["sheet"])
    assert host.distinct_rendered() <= 2


class _EditedCrops(_GalleryHost):
    """A host whose crops are changed on their way back to the replay."""

    def __init__(self, edit) -> None:
        super().__init__()
        self._edit = edit

    def chart_images(self, data: bytes, *, path: str) -> list[dict]:
        return self._edit(super().chart_images(data, path=path))


def test_out_of_range_chart_numbers_are_ignored() -> None:
    host = _EditedCrops(lambda crops: [{**crop, "chart": crop["chart"] + 5} for crop in crops])
    data = _preview(_chart_workbook(), host)

    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_a_missing_crop_still_places_the_others() -> None:
    host = _EditedCrops(lambda crops: [crop for crop in crops if crop["chart"] != 0])
    data = _preview(_two_chart_workbook(), host)

    first, second = data["workbook"]["sheet"]
    assert not first.get("images")
    assert [i["id"] for i in second["images"]] == ["chart_1"]
    assert data["chartPages"] == [2]  # the chart without a picture keeps its page listed


def test_crop_without_png_is_not_placed() -> None:
    data = _preview(_chart_workbook(), _GalleryHost(png=b""))

    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


@pytest.mark.parametrize(
    "damage",
    [
        {"png": "notbytes"},
        {"png": 5},
        {"png": b""},
        {"png": b"not PNG"},
        {"png": b"\x89PNG\r\n\x1a\n"},  # the signature and nothing else
        {"png": _png(0, 2)},
        {"png": _png(2, 0)},
        {"png": b"\x89PNG\r\n\x1a\n" + _png()[8:].replace(b"IHDR", b"IDAT", 1)},
        {"box": [50, 100, -300, 0]},
        {"box": [50, 100, 0, 200]},
        {"box": [50, 100, 300, -5]},
        {"box": [50, 100, float("inf"), 200]},
        {"box": "wide"},
        {"box": [50, 100, 300]},
        {"box": [50, 100, float("nan"), 200]},
        {"box": [50, 100, "x", 200]},
        {"salt": None},
        {"chart": "0"},
    ],
)
def test_a_damaged_crop_is_never_placed_and_never_raises(damage: dict) -> None:
    host = _EditedCrops(lambda crops: [{**crop, **damage} for crop in crops])
    data = _preview(_chart_workbook(), host)

    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_a_crop_that_is_not_a_dict_is_ignored() -> None:
    host = _EditedCrops(lambda crops: ["notadict", None, *crops[:0]])
    data = _preview(_chart_workbook(), host)

    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_a_failure_while_building_pictures_leaves_the_grid_untouched(monkeypatch) -> None:
    import langchain_canvas.replay as replay

    real = replay._chart_picture
    calls = []

    def _fails_on_the_second(number: int, crop: dict, box: dict) -> dict:
        calls.append(number)
        if len(calls) == 2:
            raise ValueError("cannot encode")
        return real(number, crop, box)

    monkeypatch.setattr(replay, "_chart_picture", _fails_on_the_second)
    data = _preview(_two_chart_workbook(), _GalleryHost())

    assert calls == [0, 1]
    assert not any(sheet.get("images") for sheet in data["workbook"]["sheet"])
    assert data["chartPages"] == [2]


def test_chartsheet_workbook_places_worksheet_charts() -> None:
    wb = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    wb.create_chartsheet("차트 시트", 0).add_chart(_bar_chart(wb.worksheets[0]))
    data = _preview(_save(wb), _GalleryHost())

    sheets = data["workbook"]["sheet"]
    assert [i["id"] for i in sheets[0]["images"]] == ["chart_0"]
    assert _box(sheets[0]["images"][0]) == (0, 280, 567, 283)
    assert data["chartPages"] == [2]  # the chartsheet chart has no anchor cell


def test_marker_crop_is_rendered_at_reading_size() -> None:
    marker = "BT /F1 6 Tf 20 800 Td (lcx0123456789abchart0) Tj ET"
    outline = "0 0 0 RG 1 w 50 300 300 200 re S"
    crops = PdfSourceConverter().chart_images(
        _vector_pdf([marker + " " + outline]), path="gallery.pdf"
    )

    width, height = Image.open(io.BytesIO(crops[0]["png"])).size
    assert 900 <= width <= 1000
    assert 560 <= height <= 700


def test_crop_with_a_foreign_salt_is_not_placed() -> None:
    host = _GalleryHost(salt="ffffffffffff")
    data = _preview(_chart_workbook(), host)

    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_chartsheet_only_workbook_keeps_chart_pages() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    _fill(ws)
    wb.create_chartsheet("차트 시트").add_chart(_bar_chart(ws))
    host = _GalleryHost()
    data = _preview(_save(wb), host)

    assert data["chartPages"] == [2]
    assert host.gallery_calls == 0


def test_hidden_sheet_chart_is_left_out_and_visible_ones_are_placed() -> None:
    wb = openpyxl.load_workbook(io.BytesIO(_two_chart_workbook()))
    wb["첫째"].sheet_state = "hidden"
    host = _GalleryHost()
    data = _preview(_save(wb), host)

    hidden, shown = data["workbook"]["sheet"]
    assert not hidden["images"]
    assert [i["id"] for i in shown["images"]] == ["chart_0"]
    assert host.titles == {0: "차트 둘째"}  # the gallery copy never held the hidden chart
    assert data["chartPages"] is None


def test_picture_bytes_stop_at_the_payload_cap(caplog) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    _fill(ws)
    for row in (15, 40, 65, 90):
        ws.add_chart(_bar_chart(ws), f"A{row}")
    host = _GalleryHost(png=_png(pad=3_000_000))

    with caplog.at_level("WARNING", logger="langchain_canvas.replay"):
        data = _preview(_save(wb), host)

    assert [i["id"] for i in data["workbook"]["sheet"][0]["images"]] == ["chart_0", "chart_1"]
    capped = [r.getMessage() for r in caplog.records if "chart_pictures_capped" in r.getMessage()]
    assert capped == ["xlsx_preview.chart_pictures_capped path=report.xlsx omitted=2"]
    assert data["chartPages"] == [2]  # the two charts left out are listed by page


def test_more_than_forty_charts_fall_back_to_chart_pages() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    _fill(ws)
    for number in range(41):
        ws.add_chart(_bar_chart(ws), f"A{15 + number * 20}")
    host = _GalleryHost()
    data = _preview(_save(wb), host)

    assert host.gallery_calls == 0
    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_a_host_that_cannot_crop_charts_gets_chart_pages() -> None:
    host = _XlsxPages()  # renders pages, has no chart_images
    store = InMemoryCanvasStore()
    commit = store.write_bytes("c1", "report.xlsx", _chart_workbook(), "chart")
    data = _file_preview_data(store, "c1", "report.xlsx", commit.revision, [host])

    assert data is not None
    assert not data["workbook"]["sheet"][0].get("images")
    assert data["chartPages"] == [2]


def test_chart_entries_report_slot_kind_and_the_unboxed_ones() -> None:
    wb = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    sheet = wb.worksheets[0]
    sheet.add_chart(
        _bar_chart(sheet), AbsoluteAnchor(pos=XDRPoint2D(0, 0), ext=XDRPositiveSize2D(9525, 9525))
    )
    wb.create_chartsheet("차트 시트").add_chart(_bar_chart(sheet))

    data = _save(wb)
    _, charts = xlsx_sheets_and_charts(data, read_only=True)

    slot = charts[0]["slot"]
    assert charts == [
        {"sheet": 0, "slot": slot, "kind": "chart", "left": 0, "top": 280, "width": 567,
         "height": 283},
        {"sheet": 0, "kind": "unboxed"},
        {"sheet": None, "kind": "unboxed"},
    ]
    # the slot is the chart's place in the reloaded worksheet's own list
    anchored = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]._charts[slot]
    assert hasattr(anchored.anchor, "_from")


# --- the existing workbook fixture rendered by a real LibreOffice --------------------


_MAC_SOFFICE = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
_SOFFICE = shutil.which("soffice") or (str(_MAC_SOFFICE) if _MAC_SOFFICE.exists() else None)


def _pdf_of(data: bytes, tmp_path: Path) -> bytes:
    source = tmp_path / "gallery.xlsx"
    source.write_bytes(data)
    out_dir = tmp_path / "out"
    out_dir.mkdir()
    subprocess.run(
        [
            str(_SOFFICE), "--headless", "--norestore", "--nolockcheck", "--nodefault",
            f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}",
            "--convert-to", "pdf", "--outdir", str(out_dir), str(source),
        ],
        check=True, capture_output=True, timeout=80,
    )
    return (out_dir / "gallery.pdf").read_bytes()


def _text_in(pdf: bytes, crop: dict) -> str:
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(pdf)
    try:
        page = document[crop["page"] - 1]
        _, page_height = page.get_size()
        x, y, width, height = crop["box"]
        return page.get_textpage().get_text_bounded(
            left=x, bottom=page_height - (y + height), right=x + width, top=page_height - y
        )
    finally:
        document.close()


@pytest.mark.skipif(_SOFFICE is None, reason="needs LibreOffice (soffice)")
def test_libreoffice_gallery_places_b04_charts(tmp_path) -> None:
    """A formatting-only first sheet used to shift every chart one sheet over."""
    data = (_FIXTURES / "formatting_first.xlsx").read_bytes()
    _, entries = xlsx_sheets_and_charts(data, read_only=True)
    boxes = [e for e in entries if e["kind"] == "chart"]
    titles = {
        sheet: [_chart_title(chart) for chart in book_sheet._charts]
        for sheet, book_sheet in enumerate(openpyxl.load_workbook(io.BytesIO(data)).worksheets)
    }

    pdf = _pdf_of(xlsx_chart_gallery(data, boxes), tmp_path)
    crops = PdfSourceConverter().chart_images(pdf, path="gallery.pdf")

    assert sorted(c["chart"] for c in crops) == list(range(len(boxes)))
    assert {c["salt"] for c in crops} == {gallery_salt(data)}
    for crop in crops:
        want = titles[boxes[crop["chart"]]["sheet"]][0]
        assert want  # the fixture titles its charts S2-C1, S3-C1
        assert want in _text_in(pdf, crop).replace("\r", " ").replace("\n", " ")
