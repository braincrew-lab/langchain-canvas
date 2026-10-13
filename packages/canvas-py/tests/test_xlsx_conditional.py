"""Conditional formats of a read-only workbook, decided against the file's values."""

from __future__ import annotations

import io
from typing import Any

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl import Workbook  # noqa: E402
from openpyxl.formatting.rule import (  # noqa: E402
    CellIsRule,
    ColorScaleRule,
    DataBarRule,
    FormulaRule,
    IconSetRule,
    Rule,
)
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.styles.differential import DifferentialStyle  # noqa: E402

from langchain_canvas.xlsx_import import xlsx_sheets_and_charts  # noqa: E402

RED = Font(color="FF0000")


def _row_workbook(values: list[Any]) -> tuple[Workbook, Any]:
    wb = Workbook()
    ws = wb.active
    for column, value in enumerate(values, start=1):
        ws.cell(row=1, column=column, value=value)
    return wb, ws


def _cells(wb: Workbook) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    out = io.BytesIO()
    wb.save(out)
    parsed, _ = xlsx_sheets_and_charts(out.getvalue(), read_only=True)
    sheet = parsed["sheets"][0]
    return {f"{c['r']}_{c['c']}": c["v"] for c in sheet["celldata"]}, sheet


def test_cell_is_less_than_zero_paints_red_font() -> None:
    wb, ws = _row_workbook([-5, 0, 5])
    ws.conditional_formatting.add("A1:C1", CellIsRule(operator="lessThan", formula=["0"], font=RED))
    cells, _ = _cells(wb)
    assert cells["0_0"]["fc"] == "#FF0000"
    assert cells["0_1"].get("fc") != "#FF0000" and cells["0_2"].get("fc") != "#FF0000"


def test_greater_or_equal_between_and_not_equal() -> None:
    wb, ws = _row_workbook([1, 5, 10])
    fill = PatternFill("solid", fgColor="FFFF00")
    ws.conditional_formatting.add(
        "A1:C1", CellIsRule(operator="greaterThanOrEqual", formula=["5"], fill=fill)
    )
    ws.conditional_formatting.add(
        "A1:C1", CellIsRule(operator="between", formula=["0", "2"], font=Font(bold=True))
    )
    ws.conditional_formatting.add(
        "A1:C1", CellIsRule(operator="notEqual", formula=["5"], font=Font(italic=True))
    )
    cells, _ = _cells(wb)
    assert (
        "bg" not in cells["0_0"]
        and cells["0_1"]["bg"] == "#FFFF00"
        and cells["0_2"]["bg"] == "#FFFF00"
    )
    assert cells["0_0"]["bl"] == 1 and "bl" not in cells["0_2"]
    assert cells["0_0"]["it"] == 1 and "it" not in cells["0_1"]


def test_priority_and_stop_if_true() -> None:
    wb, ws = _row_workbook([-5])
    first = CellIsRule(operator="lessThan", formula=["0"], font=RED, stopIfTrue=True)
    second = CellIsRule(operator="lessThan", formula=["100"], font=Font(bold=True, color="0000FF"))
    ws.conditional_formatting.add("A1", first)
    ws.conditional_formatting.add("A1", second)
    cells, _ = _cells(wb)
    assert cells["0_0"]["fc"] == "#FF0000" and "bl" not in cells["0_0"]


def test_number_format_colour_beats_a_rule_font_colour() -> None:
    wb, ws = _row_workbook([-5])
    ws["A1"].number_format = "0;[Red]-0"
    ws.conditional_formatting.add(
        "A1", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="00FF00"))
    )
    cells, _ = _cells(wb)
    assert cells["0_0"]["fc"] == "#FF0000"


def test_colour_scale_three_points_interpolate() -> None:
    wb, ws = _row_workbook([0, 50, 100])
    ws.conditional_formatting.add(
        "A1:C1",
        ColorScaleRule(
            start_type="min",
            start_color="FF0000",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFFF00",
            end_type="max",
            end_color="00FF00",
        ),
    )
    cells, _ = _cells(wb)
    assert [cells[f"0_{c}"]["bg"] for c in range(3)] == ["#FF0000", "#FFFF00", "#00FF00"]


def test_data_bar_lengths_and_icons() -> None:
    wb, ws = _row_workbook([0, 25, 100])
    ws.conditional_formatting.add(
        "A1:C1", DataBarRule(start_type="min", end_type="max", color="638EC6")
    )
    cells, _ = _cells(wb)
    assert [cells[f"0_{c}"]["lcx"]["bar"]["ratio"] for c in range(3)] == [0, 0.25, 1]

    wb, ws = _row_workbook([0, 50, 100])
    ws.conditional_formatting.add("A1:C1", IconSetRule("3Arrows", "percent", [0, 33, 67]))
    cells, _ = _cells(wb)
    assert [cells[f"0_{c}"]["lcx"]["icon"] for c in range(3)] == [
        {"set": "3Arrows", "index": 0},
        {"set": "3Arrows", "index": 1},
        {"set": "3Arrows", "index": 2},
    ]


def test_duplicates_top_and_above_average() -> None:
    wb, ws = _row_workbook([1, 2, 2, 9])
    dup = Rule(type="duplicateValues", dxf=DifferentialStyle(font=Font(bold=True)))
    ws.conditional_formatting.add("A1:D1", dup)
    top = Rule(type="top10", rank=1, dxf=DifferentialStyle(font=Font(italic=True)))
    ws.conditional_formatting.add("A1:D1", top)
    above = Rule(
        type="aboveAverage", aboveAverage=True, dxf=DifferentialStyle(font=Font(strike=True))
    )
    ws.conditional_formatting.add("A1:D1", above)
    cells, _ = _cells(wb)
    assert [cells[f"0_{c}"].get("bl") for c in range(4)] == [None, 1, 1, None]
    assert [cells[f"0_{c}"].get("it") for c in range(4)] == [None, None, None, 1]
    assert [cells[f"0_{c}"].get("cl") for c in range(4)] == [None, None, None, 1]


def test_expression_moves_relative_references() -> None:
    wb = Workbook()
    ws = wb.active
    for row, (a, b) in enumerate([(1, 2), (5, 2)], start=1):
        ws.cell(row=row, column=1, value=a)
        ws.cell(row=row, column=2, value=b)
    ws.conditional_formatting.add(
        "A1:A2", FormulaRule(formula=["AND($A1<$B1,ISBLANK(C1))"], font=RED)
    )
    cells, _ = _cells(wb)
    assert cells["0_0"]["fc"] == "#FF0000"
    assert cells["1_0"].get("fc") != "#FF0000"


def test_an_expression_outside_the_evaluator_is_listed_not_dropped() -> None:
    wb, ws = _row_workbook([1])
    ws.conditional_formatting.add("A1", FormulaRule(formula=["VLOOKUP(A1,Z1:Z9,1,0)>0"], font=RED))
    cells, sheet = _cells(wb)
    assert cells["0_0"].get("fc") != "#FF0000"
    assert sheet["lcxSkipped"][0]["type"] == "expression"


def test_a_rule_fill_covers_the_cell_pattern() -> None:
    wb, ws = _row_workbook([5])
    ws["A1"].fill = PatternFill("lightDown", fgColor="000000", bgColor="FFFFFF")
    ws.conditional_formatting.add(
        "A1",
        CellIsRule(
            operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="00FF00")
        ),
    )
    cells, _ = _cells(wb)
    assert cells["0_0"]["bg"] == "#00FF00" and "fill" not in cells["0_0"].get("lcx", {})
