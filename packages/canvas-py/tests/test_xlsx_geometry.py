"""Where a drawing anchor lands in the grid: sized, folded and far-away lines.

``_anchor_box`` answers from a sheet's sized and folded lines only, so an anchor at the
last row of a sheet costs no more than one at the first. These tests hold its answers to
what the earlier range-sum geometry gave for the same anchors (the literals below were
taken from that code before it was replaced) and bound its cost.
"""

from __future__ import annotations

import io
import time
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from PIL import Image
from test_xlsx_gallery import _bar_chart, _save

_FIXTURES = Path(__file__).parent / "fixtures"


def _geometry_workbook() -> bytes:
    """Sized columns and rows, folded rows and columns, charts and pictures."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "geo"
    for row in range(1, 13):
        ws.append([f"항목 {row}", row * 1000, "비고 설명 글자가 긴 칸", row])
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["B"].hidden = True
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[5].height = 45
    ws.row_dimensions[3].hidden = True
    ws.add_chart(_bar_chart(ws, "one cell"), "E10")
    spanning = _bar_chart(ws, "two cell")
    spanning.anchor = TwoCellAnchor(
        _from=AnchorMarker(col=3, row=12), to=AnchorMarker(col=8, row=20)
    )
    ws.add_chart(spanning)
    for cell in ("F4", "C25"):
        png = io.BytesIO()
        Image.new("RGB", (40, 30), "red").save(png, format="PNG")
        png.seek(0)
        ws.add_image(XlImage(png), cell)
    return _save(wb)


def _marker(col: int, row: int, col_off: int = 0, row_off: int = 0) -> SimpleNamespace:
    return SimpleNamespace(col=col, row=row, colOff=col_off, rowOff=row_off)


# What the range-sum geometry answered for these anchors on a sheet whose
# columns 0/2/5 are 120/80/200px, rows 1/4/10 are 40/60/25px, rows 2 and 7 and column 1
# are folded (border 0, then border 1).
_OLD_BOXES = {
    0: {
        "one_cell_origin": {"left": 0, "top": 0, "width": 100, "height": 50},
        "one_cell_offset": {"left": 300, "top": 177, "width": 567, "height": 283},
        "two_cell": {"left": 120, "top": 239, "width": 426, "height": 171},
        "starts_on_hidden": {"left": 120, "top": 59, "width": 64, "height": 32},
        "two_cell_hidden_end": {"left": 0, "top": 59, "width": 692, "height": 136},
        "far": {"left": 2299, "top": 7632, "width": 200, "height": 100},
    },
    1: {
        "one_cell_origin": {"left": 0, "top": 0, "width": 100, "height": 50},
        "one_cell_offset": {"left": 302, "top": 182, "width": 567, "height": 283},
        "two_cell": {"left": 121, "top": 248, "width": 430, "height": 180},
        "starts_on_hidden": {"left": 121, "top": 61, "width": 64, "height": 32},
        "two_cell_hidden_end": {"left": 0, "top": 61, "width": 699, "height": 141},
        "far": {"left": 2328, "top": 8030, "width": 200, "height": 100},
    },
}


def _anchors() -> dict[str, SimpleNamespace]:
    return {
        "one_cell_origin": SimpleNamespace(
            _from=_marker(0, 0), to=None, ext=SimpleNamespace(cx=100 * 9525, cy=50 * 9525)
        ),
        "one_cell_offset": SimpleNamespace(
            _from=_marker(3, 6, 952500, 190500), to=None,
            ext=SimpleNamespace(cx=5400000, cy=2700000),
        ),
        "two_cell": SimpleNamespace(_from=_marker(2, 11), to=_marker(6, 20), ext=None),
        "starts_on_hidden": SimpleNamespace(
            _from=_marker(1, 2), to=None, ext=SimpleNamespace(cx=64 * 9525, cy=32 * 9525)
        ),
        "two_cell_hidden_end": SimpleNamespace(_from=_marker(0, 3), to=_marker(8, 9), ext=None),
        "far": SimpleNamespace(
            _from=_marker(30, 400, 9525, 19050), to=None,
            ext=SimpleNamespace(cx=200 * 9525, cy=100 * 9525),
        ),
    }


def _geometry(border: int):
    from langchain_canvas.xlsx_import import _SheetGeometry

    return _SheetGeometry(
        {0: 120, 2: 80, 5: 200},
        {1: 40, 4: 60, 10: 25},
        border=border,
        hidden=(frozenset({2, 7}), frozenset({1})),
    )


@pytest.mark.parametrize("border", [0, 1])
def test_geometry_answers_like_the_range_sum_it_replaces(border: int) -> None:
    from langchain_canvas.xlsx_import import _anchor_box

    geometry = _geometry(border)
    for name, anchor in _anchors().items():
        assert _anchor_box(anchor, geometry) == _OLD_BOXES[border][name], name


def test_huge_anchor_is_clamped() -> None:
    from langchain_canvas.xlsx_import import _anchor_box

    far = SimpleNamespace(
        _from=_marker(10**9, 10**9), to=None, ext=SimpleNamespace(cx=9525, cy=9525)
    )
    started = time.perf_counter()
    box = _anchor_box(far, _geometry(0))
    assert time.perf_counter() - started < 1
    # Excel's last cell is column 16383, row 1048575 (0-based): 73px columns and 19px
    # rows, less the folded ones, plus what the sized ones add over the default.
    assert box is not None
    assert (box["left"], box["top"]) == (16382 * 73 + 47 + 7 + 127, 1048573 * 19 + 21 + 41 + 6)


def test_five_hundred_max_row_anchors_finish_fast() -> None:
    from langchain_canvas.xlsx_import import _anchor_box

    geometry = _geometry(1)
    last_row = SimpleNamespace(
        _from=_marker(3, 1048575), to=None, ext=SimpleNamespace(cx=9525, cy=9525)
    )
    started = time.perf_counter()
    boxes = [_anchor_box(last_row, geometry) for _ in range(500)]
    assert time.perf_counter() - started < 2
    assert all(box == boxes[0] for box in boxes)
    assert boxes[0] is not None and boxes[0]["top"] == 1048573 * 20 + 21 + 41 + 6


def test_five_hundred_sheets_with_a_max_row_anchor_finish_fast() -> None:
    from langchain_canvas.xlsx_import import _anchor_box, _SheetGeometry

    last_row = SimpleNamespace(
        _from=_marker(0, 1048575), to=None, ext=SimpleNamespace(cx=9525, cy=9525)
    )
    started = time.perf_counter()
    for _ in range(500):
        geometry = _SheetGeometry({0: 120}, {1: 40}, border=1)
        assert _anchor_box(last_row, geometry) is not None
    assert time.perf_counter() - started < 2


def test_anchors_past_the_workbook_budget_are_left_unplaced(monkeypatch) -> None:
    import langchain_canvas.xlsx_import as xlsx_import

    monkeypatch.setattr(xlsx_import, "_MAX_BOXED_ANCHORS", 2)
    # the sized-and-folded workbook holds two charts and two pictures: pictures come
    # first on a sheet, so the budget goes to them and the charts have no box
    parsed, charts = xlsx_import.xlsx_sheets_and_charts(_geometry_workbook())

    assert len(parsed["sheets"][0]["images"]) == 2
    assert charts == [{"sheet": 0, "kind": "unboxed"}, {"sheet": 0, "kind": "unboxed"}]

    monkeypatch.setattr(xlsx_import, "_MAX_BOXED_ANCHORS", 1)
    parsed, _ = xlsx_import.xlsx_sheets_and_charts(_geometry_workbook())
    assert len(parsed["sheets"][0]["images"]) == 1


def _boxes_of(data: bytes, *, read_only: bool) -> dict:
    from langchain_canvas.xlsx_import import xlsx_sheets_and_charts

    parsed, charts = xlsx_sheets_and_charts(data, read_only=read_only)
    return {
        "charts": [
            {key: c[key] for key in ("sheet", "left", "top", "width", "height")} for c in charts
        ],
        "images": [
            [(i["left"], i["top"], i["width"], i["height"]) for i in sheet["images"]]
            for sheet in parsed["sheets"]
        ],
    }


def _same_box(sheet: int, left: int, top: int) -> dict:
    return {"sheet": sheet, "left": left, "top": top, "width": 567, "height": 283}


# Chart boxes the range-sum geometry gave four stored workbooks (tests/fixtures).
_KIT_CHART_BOXES = {
    "chart_only_sheets.xlsx": [_same_box(1, 74, 20), _same_box(2, 74, 20)],
    "two_charts_one_sheet.xlsx": [
        _same_box(0, 550, 20), _same_box(0, 1290, 20), _same_box(0, 550, 340),
        _same_box(1, 550, 20), _same_box(2, 550, 20),
    ],
    "wide_first.xlsx": [_same_box(0, 90, 680), _same_box(1, 550, 20), _same_box(2, 550, 20)],
    "manual_break.xlsx": [_same_box(0, 550, 880), _same_box(1, 550, 20), _same_box(2, 550, 20)],
}


@pytest.mark.parametrize("name", sorted(_KIT_CHART_BOXES))
def test_geometry_matches_the_old_computation_on_the_repro_books(name: str) -> None:
    for read_only in (False, True):
        found = _boxes_of((_FIXTURES / name).read_bytes(), read_only=read_only)
        assert found["charts"] == _KIT_CHART_BOXES[name]


# Pictures and charts of _geometry_workbook() under the range-sum geometry, read_only False
# then True. Only the read-only grid folds the hidden row and column, so the answers differ.
_GEOMETRY_BOOK_OLD = {
    False: {
        "charts": [
            {"sheet": 0, "left": 495, "top": 242, "width": 567, "height": 283},
            {"sheet": 0, "left": 438, "top": 302, "width": 353, "height": 160},
        ],
        "images": [[(564, 78, 40, 30), (311, 518, 40, 30)]],
    },
    True: {
        "charts": [
            {"sheet": 0, "left": 398, "top": 222, "width": 567, "height": 283},
            {"sheet": 0, "left": 341, "top": 282, "width": 353, "height": 160},
        ],
        "images": [[(468, 59, 40, 30), (215, 499, 40, 30)]],
    },
}


@pytest.mark.parametrize("read_only", [False, True])
def test_geometry_matches_the_old_computation_with_sized_and_folded_lines(
    read_only: bool,
) -> None:
    assert _boxes_of(_geometry_workbook(), read_only=read_only) == _GEOMETRY_BOOK_OLD[read_only]
