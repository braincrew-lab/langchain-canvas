"""Repository PPTX exporter — OOXML facts of the shared parity deck.

Gate (a) of ``.omb/plans/2026-09-09-office-render-parity.md``: the deck in
``canvas-react/src/export/__fixtures__/parity-deck.slides.json`` is written by
the real :class:`SlidesPptxExporter` and read back with python-pptx. The
fixture is the same file the browser gates mount, so the three surfaces are
compared on one input.

Phase 1 (task 1) holds three permanent locks of writer behaviour that must not
move, plus ONE temporary characterisation of a known defect (the diagonal
connector, F6) that task 9 deletes and replaces with
``test_the_line_element_has_no_vertical_drop``.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from langchain_canvas.exporters import SlidesPptxExporter
from langchain_canvas.slide_text import DEFAULT_LINE_HEIGHT, grown_height_pct, wrapped_lines

FIXTURE = (
    Path(__file__).resolve().parents[2]
    / "canvas-react"
    / "src"
    / "export"
    / "__fixtures__"
    / "parity-deck.slides.json"
)
EMU_PER_INCH = 914400
PAGE_H_IN = 5.625
# The 96 dpi page the estimator converges on (U1, governing decision 1):
# 10 x 5.625 in at the density ``fontSize`` is stored at. Anchored here as
# numbers, not as the estimator's own constants, so the test is RED while the
# writer still measures on the 1280 x 720 px sheet.
METRIC_PAGE_W_PX, METRIC_PAGE_H_PX = 960.0, 540.0


def _element(element_id: str) -> dict[str, Any]:
    elements = _load_deck()["data"]["slides"][0]["elements"]
    return next(el for el in elements if el["id"] == element_id)


def _load_deck() -> dict[str, Any]:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def shapes() -> dict[str, Any]:
    """The exported deck's shapes keyed by fixture element id.

    The writer emits one shape per element in element order and gives shapes
    no names, so order is the join key; the count assertion keeps that join
    honest.
    """
    envelope = _load_deck()
    elements = envelope["data"]["slides"][0]["elements"]
    result = SlidesPptxExporter().export(json.dumps(envelope), path="parity.slides.json")
    deck = Presentation(io.BytesIO(result.data))
    (slide,) = deck.slides
    found = list(slide.shapes)
    assert len(found) == len(elements), [s.shape_type for s in found]
    return {el["id"]: shape for el, shape in zip(elements, found, strict=True)}


# --- permanent locks (writer behaviour that must survive the plan) -------------


def test_lh12_explicit_line_height_is_written_as_spcPct_120000(shapes: dict[str, Any]) -> None:
    """An explicit ``lineHeight: 1.2`` round-trips as ``a:spcPct val=120000``
    on every paragraph (governing decision 3: explicit leading is preserved)."""
    paragraphs = shapes["lh12"].text_frame.paragraphs
    assert len(paragraphs) == 2
    for paragraph in paragraphs:
        pct = paragraph._p.find(".//" + qn("a:spcPct"))
        assert pct is not None, "explicit lineHeight must write a:lnSpc/a:spcPct"
        assert pct.get("val") == "120000"


def test_plain_default_leading_writes_no_lnSpc(shapes: dict[str, Any]) -> None:
    """A paragraph with no ``lineHeight`` writes no ``a:lnSpc`` at all — the
    writer keeps the viewer's single spacing for the default (F3 stays
    documented, not converted)."""
    (paragraph,) = shapes["plain"].text_frame.paragraphs
    assert paragraph._p.find(".//" + qn("a:lnSpc")) is None


def test_rect_is_written_as_a_preset_rectangle(shapes: dict[str, Any]) -> None:
    """The file's rectangle is ``prst="rect"`` — square corners — which is
    what U4 makes the browser surfaces draw too (F5)."""
    geometry = shapes["rect"]._element.find(".//" + qn("a:prstGeom"))
    assert geometry is not None
    assert geometry.get("prst") == "rect"


# --- U1: the file is written on the 96 dpi page (task 3) --------------------------


def test_a_growing_box_is_written_at_its_96dpi_height(shapes: dict[str, Any]) -> None:
    """U1 test 6 — the ``grow`` box's written ``a:ext`` height is the height its
    lines need on the 96 dpi page (percent of 5.625 in), within 0.05 %, and the
    estimator's own ``grown_height_pct`` says the same number."""
    element = _element("grow")
    text, size, w, h = element["text"], element["fontSize"], element["w"], element["h"]
    needed_px = wrapped_lines(text, size, w / 100.0 * METRIC_PAGE_W_PX) * size * DEFAULT_LINE_HEIGHT
    expected_pct = max(h, needed_px / METRIC_PAGE_H_PX * 100.0)  # 3 lines -> 16.0
    written_pct = int(shapes["grow"].height) / EMU_PER_INCH * 100.0 / PAGE_H_IN
    assert abs(written_pct - expected_pct) <= 0.05, (written_pct, expected_pct)
    assert abs(grown_height_pct(text, size, w, h) - expected_pct) <= 0.001


@pytest.mark.parametrize("element_id", ["fit_latin", "fit_ko"])
def test_a_shrinking_box_is_written_with_the_96dpi_scale(
    shapes: dict[str, Any], element_id: str
) -> None:
    """U1 test 6b — both shrink-to-fit boxes need 3 lines (86.4 px) in a
    43.2 px box on the 96 dpi page, so the writer declares
    ``a:normAutofit fontScale="50000"``."""
    body = shapes[element_id].text_frame._bodyPr
    autofit = body.find(qn("a:normAutofit"))
    assert autofit is not None, element_id
    assert autofit.get("fontScale") == "50000", (element_id, autofit.attrib)


# --- U3: the bullet hang in the file equals the browser hang (task 7) -----------


def test_the_bullet_hang_in_the_file_equals_the_browser_hang(shapes: dict[str, Any]) -> None:
    """U3 test 5 — ``grow``'s paragraph carries the list bullet the browser
    draws (``a:buChar •``) and the same 1.2 em hanging indent every browser
    surface hangs its body by: ``marL == -indent == 18 pt x 1.2 x 12700``
    (24 px -> 18 pt), the writer's ``BULLET_HANG_EM`` twin."""
    from langchain_canvas.slide_text import BULLET_HANG_EM, BULLET_PREFIX

    (paragraph,) = shapes["grow"].text_frame.paragraphs
    properties = paragraph._p.find(qn("a:pPr"))
    assert properties is not None, "the bullet paragraph writes no a:pPr"
    hang = int(24 * 0.75 * BULLET_HANG_EM * 12700)
    assert properties.get("marL") == str(hang)
    assert properties.get("indent") == f"-{hang}"
    bullet = properties.find(qn("a:buChar"))
    assert bullet is not None and bullet.get("char") == BULLET_PREFIX.strip() == "•"
    # The prefix is stripped exactly once: the run holds the body only.
    assert paragraph.runs[0].text == _element("grow")["text"][len(BULLET_PREFIX):]


# --- U4: the line element is a flat connector as thick as its box (task 9) --------


def test_the_line_element_has_no_vertical_drop(shapes: dict[str, Any]) -> None:
    """U4 test 4 — the fixture's ``line`` (``{w: 40, h: 2}``, the editor's
    default) is written as a horizontal connector on its box's centre line,
    as thick as the box: ``0.02 x 5.625 in x 72 = 8.1 pt`` (rounded to absorb
    the EMU conversion residue, as U4 test 3 does). Replaces the Phase 1
    characterisation of the diagonal defect (F6)."""
    connector = shapes["line"]
    assert int(connector.begin_y) == int(connector.end_y)
    assert int(connector.end_x) > int(connector.begin_x)
    assert round(connector.line.width.pt, 2) == 8.1
