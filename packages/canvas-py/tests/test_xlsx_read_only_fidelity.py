"""A read-only workbook carries what makes it look like Excel.

``xlsx_sheets_and_charts(data, read_only=True)`` is the reader behind a
workbook the canvas only shows (a published file, an upload a host keeps
read-only). Hidden sheets, rows and columns stay hidden; number-format
sections pick the text and colour a value gets; strike-through, rotation,
indent, patterned and gradient fills, notes, links, list validation and tab
colours come across. The editable reader (``xlsx_to_sheets``) is unchanged —
``test_xlsx_parity.py`` pins it to its golden file.
"""

from __future__ import annotations

import io
from typing import Any

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Font, GradientFill, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from langchain_canvas.xlsx_fidelity import format_number_sections  # noqa: E402
from langchain_canvas.xlsx_import import xlsx_sheets_and_charts, xlsx_to_sheets  # noqa: E402


def _bytes(workbook: Workbook) -> bytes:
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _read_only(workbook: Workbook) -> list[dict[str, Any]]:
    parsed, _charts = xlsx_sheets_and_charts(_bytes(workbook), read_only=True)
    return parsed["sheets"]


def _cells(sheet: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {f"{c['r']}_{c['c']}": c["v"] for c in sheet["celldata"]}


# --- hidden sheets, rows, columns, tab colour -------------------------------


def test_hidden_sheets_carry_hide_and_keep_their_index() -> None:
    wb = Workbook()
    wb.active.title = "a"
    wb.create_sheet("hidden").sheet_state = "hidden"
    wb.create_sheet("very").sheet_state = "veryHidden"
    wb.create_sheet("b")
    sheets = _read_only(wb)
    assert [s["name"] for s in sheets] == ["a", "hidden", "very", "b"]
    assert [s.get("hide") for s in sheets] == [None, 1, 1, None]
    assert [s["status"] for s in sheets] == [1, 0, 0, 0]


def test_first_visible_sheet_opens_when_the_first_is_hidden() -> None:
    wb = Workbook()
    wb.active.title = "data"
    wb.create_sheet("shown")
    wb["data"].sheet_state = "hidden"
    wb.active = 1
    sheets = _read_only(wb)
    assert [s["status"] for s in sheets] == [0, 1]


def test_hidden_rows_and_columns_reach_the_grid_config() -> None:
    wb = Workbook()
    ws = wb.active
    for row in range(1, 6):
        for column in range(1, 6):
            ws.cell(row=row, column=column, value=row * 10 + column)
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["B"].hidden = True
    ws.column_dimensions.group("D", "E", hidden=True)
    config = _read_only(wb)[0]["config"]
    assert config["rowhidden"] == {2: 0}
    assert config["colhidden"] == {1: 0, 3: 0, 4: 0}


def test_tab_colour() -> None:
    wb = Workbook()
    wb.active.sheet_properties.tabColor = "C00000"
    assert _read_only(wb)[0]["color"] == "#C00000"


# --- number-format sections ---------------------------------------------------


@pytest.mark.parametrize(
    ("value", "fmt", "text", "colour"),
    [
        (1500000, "#,##0;[Red]\\(#,##0\\);\\-", "1,500,000", None),
        (-1500000, "#,##0;[Red]\\(#,##0\\);\\-", "(1,500,000)", "#FF0000"),
        (0, "#,##0;[Red]\\(#,##0\\);\\-", "-", None),
        (-1500000, "#,##0;[RED]\\(#,##0\\);\\-", "(1,500,000)", "#FF0000"),
        (1500000, "[Blue]#,##0;[Red]-#,##0", "1,500,000", "#0000FF"),
        (-1500000, "[Blue]#,##0;[Red]-#,##0", "-1,500,000", "#FF0000"),
        (-0.15, "0.0%;[Red]-0.0%", "-15.0%", "#FF0000"),
        (-1500000, '#,##0"원";[Red]-#,##0"원"', "-1,500,000원", "#FF0000"),
        (1500000, "[>=1000000][Blue]#,##0;[Red]#,##0", "1,500,000", "#0000FF"),
        (-1500000, "[>=1000000][Blue]#,##0;[Red]#,##0", "1,500,000", "#FF0000"),
        (-1234.5, "#,##0.00", "-1,234.50", None),
        (1234.5, "_(#,##0.00_);[Red]\\(#,##0.00\\)", " 1,234.50 ", None),
        (-3, "[Color10]0;[Color3]-0", "-3", "#FF0000"),
    ],
)
def test_format_number_sections(value: float, fmt: str, text: str, colour: str | None) -> None:
    assert format_number_sections(value, fmt) == (text, colour)


def test_date_formats_are_left_to_the_date_formatter() -> None:
    assert format_number_sections(45000, "yyyy-mm-dd;@") is None


def test_negative_section_reaches_the_cell_and_formula_cells_get_a_type() -> None:
    wb = Workbook()
    ws = wb.active
    fmt = "#,##0;[Red]\\(#,##0\\);\\-"
    ws["A1"] = -1500000
    ws["A1"].number_format = fmt
    ws["A2"] = "=A1"
    ws["A2"].number_format = fmt
    data = _bytes(wb)
    # openpyxl writes formulas without a cached result; give A2 one.
    data = _with_cached_value(data, "A2", "-1500000")
    parsed, _ = xlsx_sheets_and_charts(data, read_only=True)
    cells = _cells(parsed["sheets"][0])
    assert cells["0_0"]["m"] == "(1,500,000)" and cells["0_0"]["fc"] == "#FF0000"
    assert cells["1_0"]["m"] == "(1,500,000)" and cells["1_0"]["fc"] == "#FF0000"
    assert cells["1_0"]["ct"]["t"] == "n"


def test_editable_reader_keeps_its_shape() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = -1500000
    ws["A1"].number_format = "#,##0;[Red]\\(#,##0\\);\\-"
    ws.sheet_properties.tabColor = "C00000"
    wb.create_sheet("h").sheet_state = "hidden"
    sheets = xlsx_to_sheets(_bytes(wb))["sheets"]
    cell = _cells(sheets[0])["0_0"]
    assert cell["m"] == "-1,500,000" and cell.get("fc") != "#FF0000"
    assert "hide" not in sheets[1] and "color" not in sheets[0]


# --- cell decorations ---------------------------------------------------------


def test_strike_rotation_and_indent() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"], ws["A4"], ws["A5"] = "s", "up", "down", "stack", "in"
    ws["A1"].font = Font(strike=True)
    ws["A2"].alignment = Alignment(textRotation=30)
    ws["A3"].alignment = Alignment(textRotation=120)
    ws["A4"].alignment = Alignment(textRotation=255)
    ws["A5"].alignment = Alignment(horizontal="left", indent=2)
    cells = _cells(_read_only(wb)[0])
    assert cells["0_0"]["cl"] == 1
    assert cells["1_0"]["rt"] == 30
    assert cells["2_0"]["rt"] == 120  # Excel and Fortune share the 91-180 = downward encoding
    assert cells["3_0"]["tr"] == "3"
    assert cells["4_0"]["lcx"]["indent"] == 2


def test_pattern_and_gradient_fills() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"] = "p", "g"
    ws["A1"].fill = PatternFill("lightDown", fgColor="4472C4", bgColor="FFFFFF")
    ws["A2"].fill = GradientFill(degree=90, stop=("FFF2CC", "ED7D31"))
    cells = _cells(_read_only(wb)[0])
    assert cells["0_0"]["lcx"]["fill"] == {"pattern": "lightDown", "fg": "#4472C4", "bg": "#FFFFFF"}
    assert cells["0_0"]["bg"]  # a blended fallback for a grid without the painter
    gradient = cells["1_0"]["lcx"]["fill"]
    assert gradient["gradient"] == "linear" and gradient["degree"] == 90
    assert gradient["stops"] == [[0, "#FFF2CC"], [1, "#ED7D31"]]
    assert cells["1_0"]["bg"] == "#FFF2CC"


def test_comment_and_hyperlinks() -> None:
    wb = Workbook()
    ws = wb.active
    wb.create_sheet("target")
    ws["A1"] = "note"
    ws["A1"].comment = Comment("check Q3", "me")
    ws["A2"] = "web"
    ws["A2"].hyperlink = "https://example.com/a"
    ws["A3"] = "inside"
    ws["A3"].hyperlink = "#'target'!B2"
    ws["A4"] = "bad"
    ws["A4"].hyperlink = "javascript:alert(1)"
    sheet = _read_only(wb)[0]
    cells = _cells(sheet)
    assert cells["0_0"]["ps"]["value"] == "check Q3"
    assert sheet["hyperlink"] == {
        "1_0": {"linkType": "webpage", "linkAddress": "https://example.com/a"},
        "2_0": {"linkType": "cellrange", "linkAddress": "target!B2"},
    }
    assert cells["1_0"]["lcx"]["link"] == "https://example.com/a"
    assert "lcx" not in cells["3_0"]


def test_list_validation_marks_cells_holding_a_value() -> None:
    wb = Workbook()
    ws = wb.active
    dv = DataValidation(type="list", formula1='"a,b"')
    ws.add_data_validation(dv)
    dv.add("B2:B3")
    ws["B2"] = "a"
    cells = _cells(_read_only(wb)[0])
    assert cells["1_1"]["lcx"]["dropdown"] is True
    assert "2_1" not in cells  # an empty cell under the rule gets no arrow


def _with_cached_value(data: bytes, ref: str, value: str) -> bytes:
    import re
    import zipfile

    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = body.decode()
                text = re.sub(
                    rf'(<c r="{ref}"[^>]*>)<f>([^<]*)</f>(?:<v\s*/>|<v></v>)?(</c>)',
                    rf"\1<f>\2</f><v>{value}</v>\3",
                    text,
                )
                body = text.encode()
            dst.writestr(item, body)
    return out.getvalue()


def test_rotated_text_grows_an_auto_height_row_and_numbers_align_right() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["A2"], ws["A3"] = "보기 글자", "보기 글자", 0
    ws["A1"].alignment = Alignment(textRotation=30)
    ws["A2"].alignment = Alignment(textRotation=30)
    ws.row_dimensions[2].height = 15  # a stored height is kept
    ws["A3"].number_format = "#,##0;-#,##0;\\-"
    sheet = _read_only(wb)[0]
    rowlen = sheet["config"]["rowlen"]
    assert rowlen[0] > 25
    assert rowlen[1] == 20  # 15pt stored, untouched
    assert _cells(sheet)["2_0"]["m"] == "-" and _cells(sheet)["2_0"]["ht"] == 2


def test_a_picture_below_hidden_rows_moves_up_with_the_grid() -> None:
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    def build(hide: bool) -> bytes:
        wb = Workbook()
        ws = wb.active
        buffer = io.BytesIO()
        PILImage.new("RGB", (10, 10), "red").save(buffer, format="PNG")
        buffer.seek(0)
        ws.add_image(XLImage(buffer), "A6")
        if hide:
            ws.row_dimensions[2].hidden = True
            ws.row_dimensions[3].hidden = True
        return _bytes(wb)

    def top(data: bytes, read_only: bool) -> int:
        parsed, _ = xlsx_sheets_and_charts(data, read_only=read_only)
        return parsed["sheets"][0]["images"][0]["top"]

    shown, folded = top(build(False), True), top(build(True), True)
    assert shown - folded == 2 * 19
    # the editable reader keeps its old geometry
    assert top(build(True), False) == shown


def test_fortune_red_marker_is_dropped_once_the_colour_is_baked() -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = -5
    ws["A1"].number_format = "[Red]0;-0"
    cell = _cells(_read_only(wb)[0])["0_0"]
    assert cell["m"] == "-5" and cell.get("fc") != "#FF0000"
    assert "[Red]" not in cell["ct"]["fa"]


def test_pictures_and_charts_use_the_rows_after_rotated_text_grows_them() -> None:
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "보기 글자 보기 글자"
    ws["A1"].alignment = Alignment(textRotation=45)
    for row in range(2, 6):
        ws.cell(row=row, column=2, value=row)
    buffer = io.BytesIO()
    PILImage.new("RGB", (10, 10), "red").save(buffer, format="PNG")
    buffer.seek(0)
    ws.add_image(XLImage(buffer), "D6")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=5))
    ws.add_chart(chart, "D6")
    parsed, charts = xlsx_sheets_and_charts(_bytes(wb), read_only=True)
    sheet = parsed["sheets"][0]
    grown = sheet["config"]["rowlen"][0]
    assert grown > 19
    # rows 1-5 above the anchor: the grown first row plus four default rows
    assert sheet["images"][0]["top"] == grown + 4 * 19
    assert charts[0]["top"] == (grown + 1) + 4 * (19 + 1)  # charts count the 1px row border
