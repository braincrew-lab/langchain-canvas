"""The drawing reader: shapes, connectors, chartex frames and charts openpyxl cannot see.

openpyxl keeps charts and pictures and drops the rest of a worksheet's drawing, so the
reader walks ``xl/drawings`` itself. It trusts neither the zip nor the XML: the archive
limits apply, odd zip names and oversized parts are refused, and a part with a DOCTYPE is
never parsed.
"""

from __future__ import annotations

import io
import os
import stat
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from langchain_canvas.converters import UnsafeArchiveError
from langchain_canvas.xlsx_drawings import (
    MAX_EMU,
    MAX_NAME_CHARS,
    MAX_OBJECT_CHARS,
    MAX_OBJECTS,
    MAX_PART_BYTES,
    MAX_TOTAL_CHARS,
    MAX_TOTAL_PART_BYTES,
    model_safe_text,
    read_drawing_objects,
    resolve_part_target,
)

_FIXTURES = Path(__file__).parent / "fixtures"

_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


# --- builders ---------------------------------------------------------------------


def _wsdr(*anchors: str) -> str:
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><xdr:wsDr xmlns:xdr="{_XDR}" xmlns:a="{_A}" '
        f'xmlns:c="{_C}" xmlns:r="{_R}" xmlns:mc="{_MC}">' + "".join(anchors) + "</xdr:wsDr>"
    )


def _marker(col: int, row: int, col_off: int = 0, row_off: int = 0) -> str:
    return (
        f"<xdr:col>{col}</xdr:col><xdr:colOff>{col_off}</xdr:colOff>"
        f"<xdr:row>{row}</xdr:row><xdr:rowOff>{row_off}</xdr:rowOff>"
    )


def _one_cell(
    body: str, *, col: int = 0, row: int = 0, col_off: int = 0, row_off: int = 0,
    cx: int = 9525, cy: int = 9525,
) -> str:
    return (
        f"<xdr:oneCellAnchor><xdr:from>{_marker(col, row, col_off, row_off)}</xdr:from>"
        f'<xdr:ext cx="{cx}" cy="{cy}"/>{body}<xdr:clientData/></xdr:oneCellAnchor>'
    )


def _sp(name: str = "S", text: str | None = "hello", *, hidden: str | None = None) -> str:
    flag = f' hidden="{hidden}"' if hidden is not None else ""
    paragraphs = "".join(
        f"<a:p><a:r><a:t>{part}</a:t></a:r></a:p>" for part in (text or "").split("\n")
    )
    body = ""
    if text is not None:
        body = f"<xdr:txBody><a:bodyPr/><a:lstStyle/>{paragraphs}</xdr:txBody>"
    return (
        f'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="{name}"{flag}/>'
        f"<xdr:cNvSpPr/></xdr:nvSpPr><xdr:spPr/>{body}</xdr:sp>"
    )


def _grp(name: str, *children: str, hidden: str | None = None) -> str:
    flag = f' hidden="{hidden}"' if hidden is not None else ""
    return (
        f'<xdr:grpSp><xdr:nvGrpSpPr><xdr:cNvPr id="9" name="{name}"{flag}/><xdr:cNvGrpSpPr/>'
        "</xdr:nvGrpSpPr><xdr:grpSpPr/>" + "".join(children) + "</xdr:grpSp>"
    )


def _chart_frame(name: str = "Chart 1") -> str:
    return (
        f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="4" name="{name}"/>'
        "<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm/>"
        f'<a:graphic><a:graphicData uri="{_C}"><c:chart r:id="rId1"/></a:graphicData></a:graphic>'
        "</xdr:graphicFrame>"
    )


def _choice(*children: str) -> str:
    return (
        f"<mc:AlternateContent><mc:Choice Requires=\"c14\">{''.join(children)}</mc:Choice>"
        f"<mc:Fallback>{_sp('FALLBACK', 'fallback text')}</mc:Fallback></mc:AlternateContent>"
    )


def _package(
    drawings: list[str | bytes | None], *, chartsheet_first: bool = False,
    absolute: bool = False, extra: dict[str, bytes] | None = None,
    compression: int = zipfile.ZIP_DEFLATED,
) -> bytes:
    """A workbook with one worksheet per entry of ``drawings`` (None: a sheet without one)."""
    prefix = "/xl/" if absolute else ""
    sheet_entries = []
    rels = []
    parts: dict[str, str | bytes] = {}
    if chartsheet_first:
        sheet_entries.append('<sheet name="Charts" sheetId="9" r:id="rId90"/>')
        rels.append(
            f'<Relationship Id="rId90" Type="{_R}/chartsheet" '
            f'Target="{prefix}chartsheets/sheet1.xml"/>'
        )
        parts["xl/chartsheets/sheet1.xml"] = "<chartsheet/>"
    for index, drawing in enumerate(drawings, 1):
        sheet_entries.append(f'<sheet name="Sheet{index}" sheetId="{index}" r:id="rId{index}"/>')
        rels.append(
            f'<Relationship Id="rId{index}" Type="{_R}/worksheet" '
            f'Target="{prefix}worksheets/sheet{index}.xml"/>'
        )
        parts[f"xl/worksheets/sheet{index}.xml"] = f'<worksheet xmlns="{_MAIN}"/>'
        if drawing is not None:
            parts[f"xl/worksheets/_rels/sheet{index}.xml.rels"] = (
                f'<Relationships xmlns="{_PKG}"><Relationship Id="rId1" Type="{_R}/drawing" '
                f'Target="../drawings/drawing{index}.xml"/></Relationships>'
            )
            parts[f"xl/drawings/drawing{index}.xml"] = drawing
    listed = "".join(sheet_entries)
    parts["xl/workbook.xml"] = (
        f'<workbook xmlns="{_MAIN}" xmlns:r="{_R}"><sheets>{listed}</sheets></workbook>'
    )
    parts["xl/_rels/workbook.xml.rels"] = (
        f'<Relationships xmlns="{_PKG}">{"".join(rels)}</Relationships>'
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression) as archive:
        for name, body in {**parts, **(extra or {})}.items():
            archive.writestr(name, body)
    return out.getvalue()


def _names(objects: list[dict]) -> list[str]:
    return [o["name"] for o in objects]


# --- what is read -------------------------------------------------------------------


def test_reads_text_boxes_rounded_rectangles_and_grouped_shapes() -> None:
    objects = read_drawing_objects((_FIXTURES / "shapes.xlsx").read_bytes())

    summary = [
        (o["sheet"], o["sheet_name"], o["kind"], o["name"], o["text"], o["cell"], o["group"])
        for o in objects
    ]
    assert summary == [
        (0, "업무흐름", "shape", "TextBox 1", "SHAPE-TEXTBOX 승인 절차 메모", "E2", None),
        (0, "업무흐름", "shape", "TextBox 2", "SHAPE-ROUNDRECT 결재 필요", "E7", None),
        (1, "흐름도", "shape", "TextBox 1", "SHAPE-GROUP-A 접수", "B3", "Group 1"),
        (1, "흐름도", "shape", "TextBox 2", "SHAPE-GROUP-B 검토", "B3", "Group 1"),
    ]
    # a grouped child sits where its group's anchor sits
    group_anchor = {
        "from": {"col": 1, "row": 2, "colOff": 0, "rowOff": 0},
        "to": {"col": 3, "row": 7, "colOff": 209550, "rowOff": 0},
        "ext": None,
    }
    assert objects[2]["anchor"] == objects[3]["anchor"] == group_anchor
    assert objects[0]["anchor"]["to"] == {"col": 8, "row": 4, "colOff": 38100, "rowOff": 0}
    assert all(o["hidden"] is False for o in objects)


def test_reads_a_chartex_frame_and_skips_its_fallback() -> None:
    objects = read_drawing_objects((_FIXTURES / "chartex_waterfall.xlsx").read_bytes())

    assert [(o["kind"], o["name"], o["text"], o["cell"]) for o in objects] == [
        ("chartex", "Chart 1", "", "D2")
    ]
    assert objects[0]["anchor"]["to"] == {"col": 10, "row": 16, "colOff": 0, "rowOff": 0}


def test_reports_a_chart_inside_a_group() -> None:
    data = _package([_wsdr(_one_cell(_grp("G", _chart_frame("In group"))))])

    objects = read_drawing_objects(data)

    assert [(o["kind"], o["name"], o["group"]) for o in objects] == [
        ("group_chart", "In group", "G")
    ]


def test_reports_a_regular_chart_inside_mc_choice() -> None:
    data = _package([_wsdr(_one_cell(_choice(_chart_frame("Choice chart"))))])

    objects = read_drawing_objects(data)

    # the Choice branch is read, the Fallback shape is not
    assert [(o["kind"], o["name"]) for o in objects] == [("choice_chart", "Choice chart")]


def test_a_top_level_regular_chart_is_left_to_openpyxl() -> None:
    data = _package([_wsdr(_one_cell(_chart_frame()))])
    assert read_drawing_objects(data) == []


def test_connectors_and_text_runs_are_read() -> None:
    connector = (
        '<xdr:cxnSp macro=""><xdr:nvCxnSpPr><xdr:cNvPr id="5" name="Arrow"/><xdr:cNvCxnSpPr/>'
        "</xdr:nvCxnSpPr><xdr:spPr/></xdr:cxnSp>"
    )
    data = _package([_wsdr(_one_cell(connector), _one_cell(_sp("Note", "first\nsecond")))])

    objects = read_drawing_objects(data)

    assert [(o["kind"], o["name"], o["text"]) for o in objects] == [
        ("connector", "Arrow", ""),
        ("shape", "Note", "first\nsecond"),
    ]


def test_chartsheets_do_not_shift_sheet_indexes() -> None:
    data = _package(
        [_wsdr(_one_cell(_sp("One"))), _wsdr(_one_cell(_sp("Two")))], chartsheet_first=True
    )

    objects = read_drawing_objects(data)

    assert [(o["sheet"], o["sheet_name"], o["name"]) for o in objects] == [
        (0, "Sheet1", "One"),
        (1, "Sheet2", "Two"),
    ]


def test_workbook_without_drawings_returns_empty_list() -> None:
    assert read_drawing_objects(_package([None, None])) == []
    assert read_drawing_objects(b"not a zip at all") == []
    book = openpyxl.Workbook()
    out = io.BytesIO()
    book.save(out)
    assert read_drawing_objects(out.getvalue()) == []


# --- what is refused ----------------------------------------------------------------


def test_entity_declaration_part_is_skipped() -> None:
    doctype = '<!DOCTYPE wsDr [<!ENTITY x "boom">]>'
    utf8 = _wsdr(_one_cell(_sp("Utf8"))).replace("?><xdr", "?>" + doctype + "<xdr", 1)
    utf16 = _wsdr(_one_cell(_sp("Utf16"))).replace("?><xdr", "?>" + doctype + "<xdr", 1)
    nul = _wsdr(_one_cell(_sp("Nul"))).encode() + b"\x00"
    data = _package(
        [utf8, utf16.encode("utf-16-le"), nul, _wsdr(_one_cell(_sp("Fine")))]
    )

    assert _names(read_drawing_objects(data)) == ["Fine"]


def test_a_bom_utf16_part_is_skipped_not_parsed() -> None:
    part = (b"\xff\xfe" + _wsdr(_one_cell(_sp("Utf16"))).encode("utf-16-le"))
    data = _package([part, _wsdr(_one_cell(_sp("Fine")))])

    assert _names(read_drawing_objects(data)) == ["Fine"]


def test_oversized_part_is_skipped_before_reading(monkeypatch) -> None:
    big = _wsdr(_one_cell(_sp("Big"))).encode() + b" " * (17 * 1024 * 1024)
    data = _package([big, _wsdr(_one_cell(_sp("Small")))], compression=zipfile.ZIP_STORED)
    opened: list[str] = []
    real_open = zipfile.ZipFile.open

    def _recording_open(self, name, *args, **kwargs):
        opened.append(name.filename if isinstance(name, zipfile.ZipInfo) else name)
        return real_open(self, name, *args, **kwargs)

    monkeypatch.setattr(zipfile.ZipFile, "open", _recording_open)

    assert _names(read_drawing_objects(data)) == ["Small"]
    assert "xl/drawings/drawing1.xml" not in opened
    assert "xl/drawings/drawing2.xml" in opened


def test_a_part_that_is_not_utf8_is_skipped() -> None:
    latin = _wsdr(_one_cell(_sp("Latin"))).encode() + b"\xe9"
    data = _package([latin, _wsdr(_one_cell(_sp("Fine")))])

    assert _names(read_drawing_objects(data)) == ["Fine"]


def test_a_part_nested_past_the_xml_depth_limit_is_skipped() -> None:
    deep = _wsdr(_one_cell(_sp("Deep") + "<a:x>" * 300 + "</a:x>" * 300))
    data = _package([deep, _wsdr(_one_cell(_sp("Fine")))])

    assert _names(read_drawing_objects(data)) == ["Fine"]


def _text_of(data: bytes, name: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(name).decode()


def test_unreadable_structural_parts_leave_no_objects() -> None:
    doctype = '<!DOCTYPE r [<!ENTITY e "v">]>'
    base = _package([_wsdr(_one_cell(_sp()))])

    def swapped(name: str, body: bytes | None) -> bytes:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w") as dst:
            for item in src.infolist():
                if item.filename == name:
                    if body is not None:
                        dst.writestr(item, body)
                    continue
                dst.writestr(item, src.read(item))
        return out.getvalue()

    rels = _text_of(base, "xl/worksheets/_rels/sheet1.xml.rels")
    workbook = _text_of(base, "xl/workbook.xml")
    assert read_drawing_objects(swapped("xl/worksheets/_rels/sheet1.xml.rels", None)) == []
    assert read_drawing_objects(
        swapped("xl/worksheets/_rels/sheet1.xml.rels", (doctype + rels).encode())
    ) == []
    assert read_drawing_objects(swapped("xl/workbook.xml", None)) == []
    assert read_drawing_objects(swapped("xl/workbook.xml", (doctype + workbook).encode())) == []


def test_objects_without_text_properties_or_frame_data_are_read_plainly() -> None:
    bare_shape = '<xdr:sp macro=""><xdr:spPr/></xdr:sp>'
    no_text = _sp("Quiet", None)
    empty_frame = '<xdr:graphicFrame macro=""><xdr:xfrm/></xdr:graphicFrame>'
    foreign = (
        '<xdr:graphicFrame macro=""><a:graphic><a:graphicData uri="urn:other"/></a:graphic>'
        "</xdr:graphicFrame>"
    )
    data = _package([_wsdr(_one_cell(bare_shape + no_text + empty_frame + foreign))])

    assert [(o["name"], o["text"]) for o in read_drawing_objects(data)] == [("", ""), ("Quiet", "")]


def test_malformed_xml_keeps_what_was_read_before_it() -> None:
    broken = _wsdr(_one_cell(_sp("Early"))).replace("</xdr:wsDr>", "<xdr:oneCellAnchor><oops>")
    data = _package([broken, _wsdr(_one_cell(_sp("Other")))])

    assert _names(read_drawing_objects(data)) == ["Early", "Other"]


def test_odd_zip_entries_make_the_whole_read_empty(tmp_path) -> None:
    base = _package([_wsdr(_one_cell(_sp()))])

    def with_entry(info: zipfile.ZipInfo, body: bytes = b"x") -> bytes:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w") as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item))
            dst.writestr(info, body)
        return out.getvalue()

    link = zipfile.ZipInfo("xl/link")
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    assert read_drawing_objects(with_entry(link)) == []
    assert read_drawing_objects(with_entry(zipfile.ZipInfo("xl/../escape.xml"))) == []
    assert read_drawing_objects(with_entry(zipfile.ZipInfo("xl\\evil.xml"))) == []
    with pytest.warns(UserWarning, match="Duplicate name"):
        duplicate = with_entry(zipfile.ZipInfo("xl/workbook.xml"))
    assert read_drawing_objects(duplicate) == []
    assert _names(read_drawing_objects(base)) == ["S"]


def test_archive_limits_still_apply() -> None:
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/blob.bin", b"\0" * 300 * 1024 * 1024)

    with pytest.raises(UnsafeArchiveError):
        read_drawing_objects(bomb.getvalue())
    with pytest.raises(UnsafeArchiveError):
        read_drawing_objects(io.BytesIO(bomb.getvalue()))


# --- relationship targets -----------------------------------------------------------

_MEMBERS = {"xl/drawings/drawing1.xml", "x.xml", "xl/worksheets/sheet1.xml"}
_SHEET = "xl/worksheets/sheet1.xml"


@pytest.mark.parametrize(
    ("source_part", "target", "expected"),
    [
        (_SHEET, "/xl/drawings/drawing1.xml", "xl/drawings/drawing1.xml"),
        (_SHEET, "../drawings/drawing1.xml", "xl/drawings/drawing1.xml"),
        (_SHEET, "../drawings/./drawing1.xml", "xl/drawings/drawing1.xml"),
        (_SHEET, "../../x.xml", "x.xml"),
        ("xl/workbook.xml", "../../x.xml", None),
        (_SHEET, "../../../x.xml", None),
        (_SHEET, "/../x.xml", None),
        (_SHEET, "..\\drawings\\drawing1.xml", None),
        (_SHEET, "a\\b.xml", None),
        (_SHEET, "../drawings/missing.xml", None),
        (_SHEET, "../drawings/drawing1.xml\x00", None),
        (_SHEET, "", None),
        (_SHEET, "//xl/drawings/drawing1.xml", None),
    ],
)
def test_relationship_targets_follow_opc_rules(source_part, target, expected) -> None:
    assert resolve_part_target(source_part, target, _MEMBERS) == expected


def test_external_relationships_are_ignored() -> None:
    data = _package([_wsdr(_one_cell(_sp("Linked")))])
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(rewritten, "w") as dst:
        for item in src.infolist():
            body = src.read(item)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                body = body.replace(b"<Relationship ", b'<Relationship TargetMode="External" ')
            dst.writestr(item, body)

    assert read_drawing_objects(rewritten.getvalue()) == []


def test_reads_an_openpyxl_made_workbook_with_drawings() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([1, 2])
    sheet.append([3, 4])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    sheet.add_chart(chart, "D2")
    out = io.BytesIO()
    book.save(out)
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as src:
        # openpyxl points its relationships at the package root: "/xl/drawings/drawing1.xml"
        sheet_rels = src.read("xl/worksheets/_rels/sheet1.xml.rels")
        assert b'Target="/xl/drawings/drawing1.xml"' in sheet_rels
        patched = io.BytesIO()
        with zipfile.ZipFile(patched, "w") as dst:
            for item in src.infolist():
                body = src.read(item)
                if item.filename == "xl/drawings/drawing1.xml":
                    body = _wsdr(_one_cell(_sp("Stamp", "approved"), col=3, row=1)).encode()
                dst.writestr(item, body)

    objects = read_drawing_objects(patched.getvalue())

    assert [(o["name"], o["text"], o["cell"]) for o in objects] == [("Stamp", "approved", "D2")]


def test_parser_is_constructed_and_refuses_on_the_runtime_python() -> None:
    # the stdlib pull parser rejects a custom TreeBuilder target on some Pythons; the
    # reader must read a normal part and skip a DOCTYPE part on the one running here
    assert len(read_drawing_objects((_FIXTURES / "shapes.xlsx").read_bytes())) == 4
    declared = '?><!DOCTYPE d [<!ENTITY e "v">]><xdr'
    doctype = _wsdr(_one_cell(_sp("X"))).replace("?><xdr", declared, 1)
    assert read_drawing_objects(_package([doctype])) == []


# --- one parse per part, bounded memory, mc:Fallback ---------------------------------


def _shared_drawing_package(sheets: int, drawing: str) -> bytes:
    """``sheets`` worksheets that all point at the one drawing part ``drawing1.xml``."""
    shared = {
        f"xl/worksheets/_rels/sheet{index}.xml.rels": (
            f'<Relationships xmlns="{_PKG}"><Relationship Id="rId1" Type="{_R}/drawing" '
            'Target="../drawings/drawing1.xml"/></Relationships>'
        )
        for index in range(2, sheets + 1)
    }
    return _package([drawing] + [None] * (sheets - 1), extra=shared)


def test_a_drawing_shared_by_many_sheets_is_parsed_once(monkeypatch) -> None:
    import langchain_canvas.xlsx_drawings as xlsx_drawings

    texts: list[str] = []
    real = xlsx_drawings._part_text

    def _spy(archive, name):
        texts.append(name)
        return real(archive, name)

    monkeypatch.setattr(xlsx_drawings, "_part_text", _spy)
    data = _shared_drawing_package(8, _wsdr(_one_cell(_sp("Shared"))))

    objects = read_drawing_objects(data)

    assert texts.count("xl/drawings/drawing1.xml") == 1
    assert [(o["sheet"], o["sheet_name"], o["name"]) for o in objects] == [
        (index, f"Sheet{index + 1}", "Shared") for index in range(8)
    ]


def test_shared_drawings_still_obey_the_object_cap() -> None:
    many = _wsdr(*(_one_cell(_sp(f"S{n}", "x")) for n in range(MAX_OBJECTS)))

    objects = read_drawing_objects(_shared_drawing_package(4, many))

    assert len(objects) == MAX_OBJECTS  # four sheets share the room of one reading


def test_shared_drawings_still_obey_the_total_text_cap() -> None:
    wordy = _wsdr(*(_one_cell(_sp(f"S{n}", "t" * MAX_OBJECT_CHARS)) for n in range(6)))

    objects = read_drawing_objects(_shared_drawing_package(3, wordy))

    assert len(objects) == MAX_TOTAL_CHARS // MAX_OBJECT_CHARS + 1


def test_the_total_of_parsed_drawing_bytes_is_budgeted(monkeypatch) -> None:
    import langchain_canvas.xlsx_drawings as xlsx_drawings

    assert MAX_TOTAL_PART_BYTES == 4 * MAX_PART_BYTES
    data = _package([_wsdr(_one_cell(_sp(name))) for name in ("A", "B", "C")])
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        size = archive.getinfo("xl/drawings/drawing1.xml").file_size
    monkeypatch.setattr(xlsx_drawings, "MAX_TOTAL_PART_BYTES", 2 * size + size // 2)

    assert _names(read_drawing_objects(data)) == ["A", "B"]  # the third would pass the budget


def test_finished_top_level_elements_are_not_kept_while_parsing(monkeypatch) -> None:
    from xml.etree import ElementTree

    widest = [0]
    original = ElementTree.XMLPullParser

    class _Spy(original):  # type: ignore[valid-type, misc]
        def read_events(self):
            root = None
            for item in super().read_events():
                root = item[1] if root is None else root
                widest[0] = max(widest[0], len(root))
                yield item

    monkeypatch.setattr(ElementTree, "XMLPullParser", _Spy)
    clutter = "<a:x/>" * 150_000
    drawing = _wsdr(clutter, _one_cell(_sp("Kept")), clutter)

    data = _package([drawing], compression=zipfile.ZIP_STORED)

    assert _names(read_drawing_objects(data)) == ["Kept"]
    # one 16 KiB chunk holds under 3,000 of these; kept, the root would hold all 300,000
    assert widest[0] < 4_000


def test_an_anchor_under_mc_fallback_is_never_read() -> None:
    alternate = (
        "<mc:AlternateContent><mc:Choice Requires=\"x\">"
        f"{_one_cell(_sp('CHOICE'))}</mc:Choice>"
        f"<mc:Fallback>{_one_cell(_sp('FALLBACK'))}</mc:Fallback></mc:AlternateContent>"
    )

    assert _names(read_drawing_objects(_package([_wsdr(alternate)]))) == ["CHOICE"]


# --- caps ---------------------------------------------------------------------------


def test_caps_bound_objects_and_text() -> None:
    many = _wsdr(*(_one_cell(_sp(f"S{n}", "x")) for n in range(MAX_OBJECTS + 100)))
    assert len(read_drawing_objects(_package([many]))) == MAX_OBJECTS

    wide = _wsdr(_one_cell(_sp("n" * 200, "t" * 3000)))
    (one,) = read_drawing_objects(_package([wide]))
    assert len(one["text"]) == MAX_OBJECT_CHARS == 2_000
    assert len(one["name"]) == MAX_NAME_CHARS == 80


def test_reading_stops_once_the_total_text_passes_the_cap() -> None:
    wordy = _wsdr(*(_one_cell(_sp(f"S{n}", "t" * MAX_OBJECT_CHARS)) for n in range(30)))

    objects = read_drawing_objects(_package([wordy]))

    assert len(objects) == MAX_TOTAL_CHARS // MAX_OBJECT_CHARS + 1


def test_out_of_range_col_row_objects_are_dropped() -> None:
    drawing = _wsdr(
        _one_cell(_sp("big col"), col=2**31),
        _one_cell(_sp("big row"), row=1_048_576),
        _one_cell(_sp("negative offset"), col_off=-1),
        _one_cell(_sp("huge offset"), row_off=MAX_EMU + 1),
        _one_cell(_sp("huge ext"), cx=MAX_EMU + 1),
        _one_cell(_sp("negative ext"), cy=-5),
        _one_cell(_sp("last cell"), col=16383, row=1_048_575),
        _one_cell(_sp("not a number"), col="x"),  # type: ignore[arg-type]
    )

    assert _names(read_drawing_objects(_package([drawing]))) == ["last cell"]


def test_absolute_and_two_cell_anchors_are_read() -> None:
    absolute = (
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="300" cy="200"/>'
        f"{_sp('Floating')}<xdr:clientData/></xdr:absoluteAnchor>"
    )
    two = (
        f"<xdr:twoCellAnchor><xdr:from>{_marker(1, 1)}</xdr:from><xdr:to>{_marker(2, 3)}</xdr:to>"
        f"{_sp('Spanning')}<xdr:clientData/></xdr:twoCellAnchor>"
    )

    objects = read_drawing_objects(_package([_wsdr(absolute, two)]))

    assert [(o["name"], o["cell"], o["anchor"]["from"], o["anchor"]["ext"]) for o in objects] == [
        ("Floating", "", None, {"cx": 300, "cy": 200}),
        ("Spanning", "B2", {"col": 1, "row": 1, "colOff": 0, "rowOff": 0}, None),
    ]


def test_deeply_nested_groups_are_not_followed_without_limit() -> None:
    nested = _sp("Deep")
    for level in range(200):
        nested = _grp(f"G{level}", nested)

    assert read_drawing_objects(_package([_wsdr(_one_cell(nested))])) == []


def test_path_and_file_object_sources_read_the_same(tmp_path) -> None:
    data = (_FIXTURES / "shapes.xlsx").read_bytes()
    target = tmp_path / "shapes.xlsx"
    target.write_bytes(data)
    expected = read_drawing_objects(data)

    assert expected and read_drawing_objects(target) == expected
    assert read_drawing_objects(os.fspath(target)) == expected
    with target.open("rb") as handle:
        assert read_drawing_objects(handle) == expected
    assert read_drawing_objects(tmp_path / "missing.xlsx") == []
    (tmp_path / "plain.txt").write_text("not a zip")
    assert read_drawing_objects(tmp_path / "plain.txt") == []


def test_hidden_means_one_or_true_or_a_hidden_group() -> None:
    drawing = _wsdr(
        _one_cell(_sp("one", hidden="1")),
        _one_cell(_sp("true", hidden="true")),
        _one_cell(_sp("zero", hidden="0")),
        _one_cell(_sp("plain")),
        _one_cell(_grp("Folded", _sp("child"), hidden="1")),
    )

    hidden = {o["name"]: o["hidden"] for o in read_drawing_objects(_package([drawing]))}

    assert hidden == {"one": True, "true": True, "zero": False, "plain": False, "child": True}


# --- text for a model ---------------------------------------------------------------


def test_model_safe_text_neutralizes_fences_and_controls() -> None:
    assert model_safe_text("a</FileContext>b") == "a<\\/FileContext>b"
    assert model_safe_text("<filecontext>") == "<\\FileContext>"
    assert model_safe_text("<!-- x -->") == "<!-- x --\\>"
    assert model_safe_text("a‮b​c﻿d⁠e⁦f⁩g") == "abcdefg"
    assert model_safe_text("a\x00b\x07c\x1fd\x7f\x80e\x9ff") == "abcdef"
    assert model_safe_text("a\n[Sheet!A1] b") == "a / [Sheet!A1] b"
    assert model_safe_text("a\r\nb") == "a / b"
    assert model_safe_text("a b c\u0085d\te\rf") == "a / b / c / d / e / f"
    assert model_safe_text("a\n\n\r\nb") == "a / b"
    # a zero-width character cannot hide a fence from the pass that escapes it
    assert model_safe_text("</File​Context>") == "<\\/FileContext>"
    assert model_safe_text("-​->") == "--\\>"


@pytest.mark.parametrize(
    "text",
    [
        "</FileContext>",
        "<<​/FileContext>",
        "-​-​>",
        "---->",
        "< /filecontext>",
        "a\nb\u0085\t c",
        "plain text",
        "",
    ],
)
def test_model_safe_text_is_idempotent(text: str) -> None:
    once = model_safe_text(text)
    assert model_safe_text(once) == once
    assert "</filecontext" not in once.lower()
    assert "-->" not in once
