"""Charts openpyxl cannot read are shown as a placeholder picture at their anchor.

An Excel 2016+ (chartex) chart, a regular chart inside a group and one inside an
``mc:Choice`` never reach the grid as charts: openpyxl drops them and LibreOffice draws
a chartex as an unrelated line chart. The drawing reader finds them, the gallery copy
puts a notice on their page, and the notice is placed where the chart was.
"""

from __future__ import annotations

import io
import re
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from test_xlsx_drawings import (
    _A,
    _C,
    _MC,
    _R,
    _XDR,
    _chart_frame,
    _choice,
    _grp,
    _marker,
    _one_cell,
)
from test_xlsx_gallery import (
    _SOFFICE,
    _bar_chart,
    _box,
    _chart_workbook,
    _EditedCrops,
    _GalleryHost,
    _pdf_of,
    _preview,
    _save,
    _text_in,
)

from langchain_canvas.converters import PdfSourceConverter
from langchain_canvas.replay import _PREVIEW_CACHE
from langchain_canvas.xlsx_gallery import (
    UNSUPPORTED_CHART_TEXT,
    gallery_salt,
    xlsx_chart_gallery,
)
from langchain_canvas.xlsx_import import xlsx_sheets_and_charts

_FIXTURES = Path(__file__).parent / "fixtures"
_CHARTEX_URI = "http://schemas.microsoft.com/office/drawing/2014/chartex"


@pytest.fixture(autouse=True)
def _clear_preview_cache():
    _PREVIEW_CACHE.clear()
    yield
    _PREVIEW_CACHE.clear()


# --- builders ---------------------------------------------------------------------


def _chartex_frame(name: str = "Waterfall", *, hidden: str | None = None) -> str:
    """A chartex frame the way Excel writes it: in mc:Choice, with a fallback shape."""
    flag = f' hidden="{hidden}"' if hidden is not None else ""
    return _choice(
        f'<xdr:graphicFrame macro=""><xdr:nvGraphicFramePr><xdr:cNvPr id="3" name="{name}"{flag}/>'
        "<xdr:cNvGraphicFramePr/></xdr:nvGraphicFramePr><xdr:xfrm/>"
        f'<a:graphic><a:graphicData uri="{_CHARTEX_URI}"><cx:chart xmlns:cx="{_CHARTEX_URI}" '
        'r:id="rId99"/></a:graphicData></a:graphic></xdr:graphicFrame>'
    )


def _declared(anchor: str) -> str:
    """Declare the prefixes a snippet uses on its anchor, so it can join any drawing."""
    names = f'xmlns:xdr="{_XDR}" xmlns:a="{_A}" xmlns:c="{_C}" xmlns:r="{_R}" xmlns:mc="{_MC}"'
    return anchor.replace(">", f" {names}>", 1)


def _two_cell(body: str, start: tuple[int, int], end: tuple[int, int]) -> str:
    return _declared(
        f"<xdr:twoCellAnchor><xdr:from>{_marker(*start)}</xdr:from><xdr:to>{_marker(*end)}</xdr:to>"
        f"{body}<xdr:clientData/></xdr:twoCellAnchor>"
    )


def _sized_one_cell(body: str, *, col: int = 3, row: int = 1, cx: int, cy: int) -> str:
    return _declared(_one_cell(body, col=col, row=row, cx=cx, cy=cy))


def _drawing_part(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rels = archive.read("xl/worksheets/_rels/sheet1.xml.rels").decode()
    found = re.search(r'Target="(?:\.\./drawings|/xl/drawings)/(drawing\d+\.xml)"', rels)
    assert found, rels
    return "xl/drawings/" + found[1]


def _with_drawing(data: bytes, edit) -> bytes:
    """``data`` with the first sheet's drawing part rewritten by ``edit(xml) -> xml``."""
    part = _drawing_part(data)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item)
            if item.filename == part:
                body = edit(body.decode()).encode()
            dst.writestr(item, body)
    return out.getvalue()


def _replacing(*anchors: str):
    """An edit that swaps the drawing's content for ``anchors`` (the chart is dropped)."""

    def edit(xml: str) -> str:
        root = re.search(r"<(?:\w+:)?wsDr\b[^>]*>", xml)
        assert root, xml
        return xml[: root.end()] + "".join(anchors) + "</" + root[0][1:].split()[0] + ">"

    return edit


def _adding(*anchors: str):
    """An edit that keeps the drawing and adds ``anchors`` after its content."""

    def edit(xml: str) -> str:
        close = re.search(r"</(?:\w+:)?wsDr>", xml)
        assert close, xml
        return xml[: close.start()] + "".join(anchors) + xml[close.start() :]

    return edit


def _placeholder_book(*anchors: str) -> bytes:
    """A one-sheet workbook whose drawing holds only ``anchors``."""
    return _with_drawing(_chart_workbook(), _replacing(*anchors))


def _entries(data: bytes) -> list[dict]:
    return xlsx_sheets_and_charts(data, read_only=True)[1]


def _kinds(entries: list[dict]) -> list[str]:
    return [entry["kind"] for entry in entries]


# --- the reader reports them ---------------------------------------------------------


def test_chartex_frame_becomes_an_unsupported_entry_in_read_only_mode() -> None:
    data = (_FIXTURES / "chartex_waterfall.xlsx").read_bytes()

    (entry,) = _entries(data)

    assert (entry["sheet"], entry["slot"], entry["kind"]) == (0, 0, "unsupported")
    assert entry["width"] > 16 and entry["height"] > 16
    assert xlsx_sheets_and_charts(data)[1] == []  # the editable reader does not look


def test_grouped_chart_becomes_an_unsupported_entry() -> None:
    data = _placeholder_book(_two_cell(_grp("G", _chart_frame()), (0, 0), (4, 6)))

    entries = _entries(data)

    # the box for the placeholder, and the chart itself listed as one with no box of its own
    assert _kinds(entries) == ["unsupported", "unboxed"]
    assert entries[1] == {"sheet": 0, "kind": "unboxed"}


def test_choice_chart_becomes_an_unsupported_entry_and_an_unboxed_one() -> None:
    data = _placeholder_book(_two_cell(_choice(_chart_frame()), (0, 0), (4, 6)))
    assert _kinds(_entries(data)) == ["unsupported", "unboxed"]


def test_a_chartex_frame_adds_no_unboxed_entry() -> None:
    data = _placeholder_book(_two_cell(_chartex_frame(), (0, 0), (4, 6)))
    assert _kinds(_entries(data)) == ["unsupported"]


def test_absolute_anchor_chartex_has_no_box_but_is_counted() -> None:
    absolute = _declared(
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="3000000" cy="2000000"/>'
        f"{_chartex_frame()}<xdr:clientData/></xdr:absoluteAnchor>"
    )
    assert _entries(_placeholder_book(absolute)) == [{"sheet": 0, "kind": "unboxed"}]


def test_one_cell_anchor_chartex_gets_its_box() -> None:
    anchor = _sized_one_cell(_chartex_frame(), cx=400 * 9525, cy=300 * 9525)

    (entry,) = _entries(_placeholder_book(anchor))

    assert (entry["width"], entry["height"]) == (400, 300)
    assert entry["kind"] == "unsupported"


def test_a_hidden_chartex_or_one_in_a_hidden_group_is_left_out() -> None:
    data = _placeholder_book(
        _two_cell(_chartex_frame(hidden="1"), (0, 0), (4, 6)),
        _two_cell(_grp("Folded", _chartex_frame(), hidden="true"), (0, 8), (4, 14)),
    )
    assert _entries(data) == []


def test_placeholders_follow_the_charts_in_top_left_order() -> None:
    low = _two_cell(_chartex_frame("Low"), (0, 30), (4, 36))
    high = _two_cell(_chartex_frame("High"), (0, 20), (4, 26))
    data = _with_drawing(_chart_workbook(), _adding(low, high))

    entries = _entries(data)

    assert _kinds(entries) == ["chart", "unsupported", "unsupported"]
    assert entries[1]["top"] < entries[2]["top"]
    assert [e["slot"] for e in entries[1:]] == [1, 0]  # drawing order: High is the second


def test_a_broken_drawing_object_is_skipped_not_fatal(monkeypatch) -> None:
    import langchain_canvas.xlsx_drawings as xlsx_drawings

    def _anchor(col: object) -> dict:
        return {
            "from": {"col": col, "row": 1, "colOff": 0, "rowOff": 0},
            "to": None, "ext": {"cx": 9525 * 200, "cy": 9525 * 100},
        }

    def _object(name: str, col: object) -> dict:
        return {
            "sheet": 0, "sheet_name": "s", "kind": "chartex", "name": name, "text": "",
            "hidden": False, "group": None, "anchor": _anchor(col), "cell": "",
        }

    monkeypatch.setattr(
        xlsx_drawings, "read_drawing_objects_checked",
        lambda source, **kwargs: ([_object("broken", "x"), _object("fine", 2)], True),
    )

    entries = _entries(_chart_workbook())

    # the broken one gets no box and is counted, the fine one and the chart are placed
    assert _kinds(entries) == ["chart", "unsupported", "unboxed"]
    assert (entries[1]["width"], entries[1]["height"]) == (200, 100)


def test_a_reader_failure_leaves_the_grid_alone(monkeypatch) -> None:
    import langchain_canvas.xlsx_drawings as xlsx_drawings

    def _boom(source, **kwargs):
        raise RuntimeError("cannot read drawings")

    monkeypatch.setattr(xlsx_drawings, "read_drawing_objects_checked", _boom)

    assert _kinds(_entries(_chart_workbook())) == ["chart"]


def test_placeholders_spend_the_anchor_budget_but_their_charts_are_still_listed(
    monkeypatch,
) -> None:
    import langchain_canvas.xlsx_import as xlsx_import

    monkeypatch.setattr(xlsx_import, "_MAX_BOXED_ANCHORS", 0)
    data = _placeholder_book(_two_cell(_grp("G", _chart_frame()), (0, 0), (4, 6)))

    assert _kinds(xlsx_import.xlsx_sheets_and_charts(data, read_only=True)[1]) == ["unboxed"]


# --- the gallery draws them ----------------------------------------------------------


def _merged_size(sheet) -> tuple[int, int]:
    (merged,) = sheet.merged_cells.ranges
    return merged.max_col - merged.min_col + 1, merged.max_row - merged.min_row + 1


def test_gallery_draws_the_unsupported_placeholder_text() -> None:
    data = (_FIXTURES / "chartex_waterfall.xlsx").read_bytes()
    entries = _entries(data)

    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, entries)))

    page = book[f"lcx{gallery_salt(data)}chart0"]
    assert page["A1"].value == f"lcx{gallery_salt(data)}chart0"
    assert page["B2"].value == UNSUPPORTED_CHART_TEXT
    assert UNSUPPORTED_CHART_TEXT == (
        "미리보기에서 지원하지 않는 차트입니다 / Chart type not supported in preview"
    )
    assert page["B2"].alignment.horizontal == "center"
    assert page["B2"].alignment.vertical == "center"
    assert page["B2"].alignment.wrap_text is True
    assert page["B2"].border.left.style == "thin" and page["B2"].border.bottom.style == "thin"
    assert book.active.title == page.title
    # about the size of the chart it stands for: 64px columns, 20px rows
    width, height = entries[0]["width"], entries[0]["height"]
    assert _merged_size(page) == (min(25, round(width / 64)), min(60, round(height / 20)))


def test_gallery_mixes_chart_and_placeholder_pages_in_entry_order() -> None:
    notice = _two_cell(_chartex_frame(), (0, 40), (4, 46))
    data = _with_drawing(_chart_workbook(), _adding(notice))
    entries = _entries(data)

    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, entries)))

    salt = gallery_salt(data)
    assert len(book[f"lcx{salt}chart0"]._charts) == 1
    assert not book[f"lcx{salt}chart1"]._charts
    assert book[f"lcx{salt}chart1"]["B2"].value == UNSUPPORTED_CHART_TEXT


def test_gallery_refuses_entries_with_nothing_to_draw() -> None:
    data = _chart_workbook()
    (box,) = _entries(data)

    with pytest.raises(ValueError):
        xlsx_chart_gallery(data, [{**box, "kind": "unboxed"}])
    # beside something to draw, an entry of another kind gets no page and keeps its number
    book = openpyxl.load_workbook(
        io.BytesIO(xlsx_chart_gallery(data, [box, {**box, "kind": "unboxed"}]))
    )
    salt = gallery_salt(data)
    assert f"lcx{salt}chart0" in book.sheetnames and f"lcx{salt}chart1" not in book.sheetnames


def test_huge_anchor_placeholder_is_bounded() -> None:
    whole_sheet = _two_cell(_chartex_frame(), (0, 0), (16383, 1_048_575))
    data = _placeholder_book(whole_sheet)
    started = time.perf_counter()

    entries = _entries(data)
    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, entries)))

    assert time.perf_counter() - started < 8
    cols, rows = _merged_size(book[f"lcx{gallery_salt(data)}chart0"])
    assert cols <= 25 and rows == 60


def test_a_wide_placeholder_stops_at_twenty_five_columns() -> None:
    anchor = _sized_one_cell(_chartex_frame(), cx=3000 * 9525, cy=100 * 9525)
    data = _placeholder_book(anchor)

    book = openpyxl.load_workbook(io.BytesIO(xlsx_chart_gallery(data, _entries(data))))

    cols, rows = _merged_size(book[f"lcx{gallery_salt(data)}chart0"])
    assert (cols, rows) == (25, 4)  # 1600px wide at most, 80px tall at least


# --- the replay places them ----------------------------------------------------------


def test_unsupported_placeholder_lands_at_its_anchor() -> None:
    data = (_FIXTURES / "chartex_waterfall.xlsx").read_bytes()
    (entry,) = _entries(data)
    host = _GalleryHost()

    preview = _preview(data, host)

    (image,) = preview["workbook"]["sheet"][0]["images"]
    assert image["id"] == "chart_0"
    assert _box(image) == (entry["left"], entry["top"], entry["width"], entry["height"])


def test_chartex_only_workbook_lists_chart_pages_only_without_its_notice(monkeypatch) -> None:
    data = (_FIXTURES / "chartex_waterfall.xlsx").read_bytes()

    placed = _preview(data, _GalleryHost())
    assert placed["workbook"]["sheet"][0]["images"] and placed["chartPages"] is None

    import langchain_canvas.xlsx_gallery as xlsx_gallery

    def _boom(data: bytes, entries: list[dict]) -> bytes:
        raise ValueError("cannot build the gallery")

    monkeypatch.setattr(xlsx_gallery, "xlsx_chart_gallery", _boom)
    _PREVIEW_CACHE.clear()
    failed = _preview(data, _GalleryHost())
    # no notice picture was placed, so the pages that carry the chart are listed
    assert not failed["workbook"]["sheet"][0]["images"] and failed["chartPages"] == [2]


def test_a_placeholder_alone_does_not_stand_for_a_missing_chart() -> None:
    data = _with_drawing(
        _chart_workbook(), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46)))
    )
    host = _EditedCrops(lambda crops: [crop for crop in crops if crop["chart"] == 1])

    preview = _preview(data, host)

    assert [i["id"] for i in preview["workbook"]["sheet"][0]["images"]] == ["chart_1"]
    assert preview["chartPages"] == [2]  # the regular chart got no picture


def test_placeholder_and_chartsheet_keep_chart_pages() -> None:
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    book.create_chartsheet("차트 시트").add_chart(_bar_chart(book.worksheets[0]))
    data = _with_drawing(
        _save(book), _replacing(_two_cell(_chartex_frame(), (0, 0), (4, 6)))
    )

    preview = _preview(data, _GalleryHost())

    assert len(preview["workbook"]["sheet"][0]["images"]) == 1  # the placeholder
    assert preview["chartPages"] == [2]  # the chartsheet chart has no anchor cell


def test_placeholder_and_absolute_anchor_chart_keep_chart_pages() -> None:
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    sheet = book.worksheets[0]
    sheet._charts.clear()
    sheet.add_chart(
        _bar_chart(sheet),
        AbsoluteAnchor(pos=XDRPoint2D(0, 0), ext=XDRPositiveSize2D(5400000, 2700000)),
    )
    data = _with_drawing(_save(book), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46))))

    preview = _preview(data, _GalleryHost())

    assert len(preview["workbook"]["sheet"][0]["images"]) == 1
    assert preview["chartPages"] == [2]


def test_placeholder_and_choice_chart_keep_chart_pages() -> None:
    data = _placeholder_book(
        _two_cell(_chartex_frame(), (0, 0), (4, 6)),
        _two_cell(_choice(_chart_frame()), (0, 10), (4, 16)),
    )

    preview = _preview(data, _GalleryHost())

    assert [i["id"] for i in preview["workbook"]["sheet"][0]["images"]] == ["chart_0", "chart_1"]
    assert preview["chartPages"] == [2]


# --- a real LibreOffice --------------------------------------------------------------


@pytest.mark.skipif(_SOFFICE is None, reason="needs LibreOffice (soffice)")
def test_libreoffice_gallery_page_of_a_chartex_carries_the_placeholder_text(tmp_path) -> None:
    data = (_FIXTURES / "chartex_waterfall.xlsx").read_bytes()
    entries = _entries(data)

    pdf = _pdf_of(xlsx_chart_gallery(data, entries), tmp_path)
    crops = PdfSourceConverter().chart_images(pdf, path="gallery.pdf")

    (crop,) = crops
    assert (crop["chart"], crop["salt"]) == (0, gallery_salt(data))
    text = _text_in(pdf, crop).replace("\r", " ").replace("\n", " ")
    assert "Chart type not supported in preview" in text
    assert "Chart Title" not in text  # what LibreOffice draws for a chartex by itself
