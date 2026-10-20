"""Every chart the preview cannot draw is accounted for, so its pages are listed.

The grid places a picture only for a chart it found and could box. A chart object that
was found but got no box (no anchor cell, past the anchor budget, an error while placing
it), and a drawing the reader had to stop reading, each count as one unboxed chart: the
rendered pages that carry charts are then listed under the grid.
"""

from __future__ import annotations

import io
import re
import zipfile

import openpyxl
import pytest
from openpyxl.drawing.image import Image
from test_xlsx_drawings import _chart_frame, _grp, _sp
from test_xlsx_drawings_caps import _pic
from test_xlsx_gallery import _bar_chart, _chart_workbook, _GalleryHost, _png, _preview, _save
from test_xlsx_placeholders import (
    _adding,
    _chartex_frame,
    _declared,
    _entries,
    _kinds,
    _two_cell,
    _with_drawing,
)

from langchain_canvas.replay import _PREVIEW_CACHE
from langchain_canvas.xlsx_drawings import MAX_OBJECT_CHARS, MAX_OBJECTS, MAX_TOTAL_CHARS


@pytest.fixture(autouse=True)
def _clear_preview_cache():
    _PREVIEW_CACHE.clear()
    yield
    _PREVIEW_CACHE.clear()


def _absolute_chartex() -> str:
    return _declared(
        '<xdr:absoluteAnchor><xdr:pos x="0" y="0"/><xdr:ext cx="3000000" cy="2000000"/>'
        f"{_chartex_frame()}<xdr:clientData/></xdr:absoluteAnchor>"
    )


def _image_ids(preview: dict) -> list[str]:
    return [image["id"] for image in preview["workbook"]["sheet"][0]["images"]]


# --- a chartex frame that got no box ----------------------------------------------------


def test_a_chartex_with_no_anchor_cell_is_counted_beside_an_ordinary_chart() -> None:
    data = _with_drawing(_chart_workbook(), _adding(_absolute_chartex()))

    assert _kinds(_entries(data)) == ["chart", "unboxed"]
    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


def test_a_chartex_past_the_anchor_budget_is_counted() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["a", 1])
    png = _png()
    for row in range(1, 500):  # 499 pictures
        sheet.add_image(Image(io.BytesIO(png)), f"D{row}")
    chart_book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    sheet.add_chart(chart_book.worksheets[0]._charts[0], "A15")
    data = _with_drawing(_save(book), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46))))

    preview = _preview(data, _GalleryHost())

    assert "chart_0" in _image_ids(preview)  # the ordinary chart took the last anchor
    assert preview["chartPages"] == [2]


def test_a_chartex_that_cannot_be_placed_is_counted_once(monkeypatch) -> None:
    import langchain_canvas.xlsx_import as xlsx_import

    def _boom(anchor):
        raise ValueError("no box")

    monkeypatch.setattr(xlsx_import, "_anchor_of", _boom)
    data = _with_drawing(_chart_workbook(), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46))))

    assert _kinds(_entries(data)) == ["chart", "unboxed"]


def test_a_placed_chartex_adds_no_unboxed_entry() -> None:
    data = _with_drawing(_chart_workbook(), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46))))

    assert _kinds(_entries(data)) == ["chart", "unsupported"]
    assert _preview(data, _GalleryHost())["chartPages"] is None


def test_a_group_chart_is_counted_once_even_when_it_has_no_box() -> None:
    data = _with_drawing(
        _chart_workbook(), _adding(_two_cell(_grp("G", _chart_frame()), (0, 40), (4, 46)))
    )

    assert _kinds(_entries(data)).count("unboxed") == 1


# --- a drawing the reader stopped reading ------------------------------------------------


def _shapes(count: int, text: str) -> list[str]:
    return [_two_cell(_sp(f"S{n}", text), (0, 100 + n), (1, 101 + n)) for n in range(count)]


def _hidden_group_chart() -> str:
    return _two_cell(_grp("G", _chart_frame()), (0, 900), (4, 906))


def test_a_grouped_chart_after_the_text_cap_still_lists_the_chart_pages() -> None:
    wordy = _shapes(MAX_TOTAL_CHARS // MAX_OBJECT_CHARS + 1, "t" * MAX_OBJECT_CHARS)
    data = _with_drawing(_chart_workbook(), _adding(*wordy, _hidden_group_chart()))

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


def test_a_grouped_chart_after_the_object_cap_still_lists_the_chart_pages() -> None:
    data = _with_drawing(
        _chart_workbook(), _adding(*_shapes(MAX_OBJECTS, ""), _hidden_group_chart())
    )

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


def test_shapes_under_every_cap_keep_the_chart_pages_unlisted() -> None:
    data = _with_drawing(_chart_workbook(), _adding(*_shapes(5, "note")))

    assert _preview(data, _GalleryHost())["chartPages"] is None


def test_a_chart_after_a_picture_at_the_object_cap_still_lists_the_chart_pages() -> None:
    members = [_sp(f"S{n}", "") for n in range(MAX_OBJECTS)]
    group = _two_cell(_grp("G", *members, _pic(), _chart_frame()), (0, 900), (4, 906))
    data = _with_drawing(_chart_workbook(), _adding(group))

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


def test_a_chart_after_a_picture_at_the_text_cap_still_lists_the_chart_pages() -> None:
    count = MAX_TOTAL_CHARS // MAX_OBJECT_CHARS + 1
    wordy = [_sp(f"S{n}", "t" * MAX_OBJECT_CHARS) for n in range(count)]
    group = _two_cell(_grp("G", *wordy, _pic(), _chart_frame()), (0, 900), (4, 906))
    data = _with_drawing(_chart_workbook(), _adding(group))

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


# --- nothing is lost after an exact cap --------------------------------------------------


def _drawing_part_of(data: bytes, sheet_number: int) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        rels = archive.read(f"xl/worksheets/_rels/sheet{sheet_number}.xml.rels").decode()
    found = re.search(r'Target="(?:\.\./drawings|/xl/drawings)/(drawing\d+\.xml)"', rels)
    assert found, rels
    return "xl/drawings/" + found[1]


def _with_part(data: bytes, part: str, edit) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            body = src.read(item)
            dst.writestr(item, edit(body.decode()).encode() if item.filename == part else body)
    return out.getvalue()


def _full_first_sheet_then(*, picture: bool) -> bytes:
    """The one-chart sheet with exactly the object cap of shapes, then a second sheet."""
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    later = book.create_sheet("Later")
    if picture:
        later.add_image(Image(io.BytesIO(_png())), "A1")
    # openpyxl drops shapes on save, so they go in after it
    return _with_drawing(_save(book), _adding(*_shapes(MAX_OBJECTS, "")))


@pytest.mark.parametrize("picture", [False, True], ids=["empty worksheet", "picture-only drawing"])
def test_an_empty_later_sheet_after_the_exact_cap_keeps_the_chart_pages_unlisted(
    picture: bool,
) -> None:
    preview = _preview(_full_first_sheet_then(picture=picture), _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] is None


def test_a_later_sheet_with_a_shape_after_the_exact_cap_lists_the_chart_pages() -> None:
    data = _full_first_sheet_then(picture=True)
    late = _two_cell(_sp("Late"), (0, 1), (1, 2))
    data = _with_part(data, _drawing_part_of(data, 2), _adding(late))

    assert _preview(data, _GalleryHost())["chartPages"] == [2]


def test_a_picture_only_anchor_after_the_exact_cap_keeps_the_chart_pages_unlisted() -> None:
    picture = _two_cell(_pic(), (0, 2), (1, 3))
    data = _with_drawing(_chart_workbook(), _adding(*_shapes(MAX_OBJECTS, ""), picture))

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] is None


# --- a hidden sheet is not printed -------------------------------------------------------


def _with_hidden_sheet(state: str, *anchors: str) -> bytes:
    """The one-chart workbook plus a second sheet in ``state`` whose drawing gains ``anchors``."""
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    hidden = book.create_sheet("Hidden")
    hidden.add_image(Image(io.BytesIO(_png())), "A1")
    hidden.sheet_state = state
    data = _save(book)
    return _with_part(data, _drawing_part_of(data, 2), _adding(*anchors))


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_an_unplaced_chartex_on_a_hidden_sheet_adds_no_unboxed_entry(state: str) -> None:
    data = _with_hidden_sheet(state, _absolute_chartex())

    preview = _preview(data, _GalleryHost())

    assert _image_ids(preview) == ["chart_0"]
    assert preview["chartPages"] is None


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_a_group_chart_on_a_hidden_sheet_adds_no_unboxed_entry(state: str) -> None:
    data = _with_hidden_sheet(state, _two_cell(_grp("G", _chart_frame()), (0, 40), (4, 46)))

    assert "unboxed" not in _kinds(_entries(data))
    assert _preview(data, _GalleryHost())["chartPages"] is None


def test_an_unplaced_chartex_on_a_visible_sheet_still_adds_an_unboxed_entry() -> None:
    data = _with_hidden_sheet("visible", _absolute_chartex())

    assert _kinds(_entries(data)).count("unboxed") == 1


def test_a_hidden_chart_sheet_adds_no_unboxed_entry() -> None:
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    sheet = book.create_chartsheet("Big")
    sheet.add_chart(_bar_chart(book.worksheets[0]))
    sheet.sheet_state = "hidden"
    data = _save(book)

    assert "unboxed" not in _kinds(_entries(data))
    assert _preview(data, _GalleryHost())["chartPages"] is None


def test_a_shown_chart_sheet_still_adds_an_unboxed_entry() -> None:
    book = openpyxl.load_workbook(io.BytesIO(_chart_workbook()))
    book.create_chartsheet("Big").add_chart(_bar_chart(book.worksheets[0]))

    assert _kinds(_entries(_save(book))).count("unboxed") == 1


def test_an_incomplete_drawing_read_still_adds_its_entry_beside_a_hidden_sheet() -> None:
    data = _with_hidden_sheet("hidden", *_shapes(MAX_OBJECTS + 1, ""))

    assert _kinds(_entries(data)) == ["chart", "unboxed"]
