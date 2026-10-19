"""Which crops of a gallery copy are placed, and how a gallery page is cropped.

A chart without a picture keeps the pages that carry it listed under the grid. A page
whose marker the converter finds is cropped around everything but that marker, even when
LibreOffice splits the marker over several text objects, and a rotated page is not cropped.
"""

from __future__ import annotations

import io

import openpyxl
import pypdfium2 as pdfium
import pytest
from test_file_preview import _vector_pdf
from test_xlsx_gallery import (
    _bar_chart,
    _EditedCrops,
    _fill,
    _GalleryHost,
    _png,
    _preview,
    _save,
)

from langchain_canvas.converters import PdfSourceConverter
from langchain_canvas.replay import _PREVIEW_CACHE


@pytest.fixture(autouse=True)
def _clear_preview_cache():
    _PREVIEW_CACHE.clear()
    yield
    _PREVIEW_CACHE.clear()


def _five_chart_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    _fill(ws)
    for row in (15, 40, 65, 90, 115):
        ws.add_chart(_bar_chart(ws), f"A{row}")
    return _save(wb)


def _ids(data: dict) -> list[str]:
    return [i["id"] for i in data["workbook"]["sheet"][0].get("images", [])]


# --- charts left without a picture keep their pages listed -----------------------------


def test_one_crop_of_five_places_one_picture_and_lists_the_chart_pages() -> None:
    host = _EditedCrops(lambda crops: crops[:1])

    data = _preview(_five_chart_workbook(), host)

    assert _ids(data) == ["chart_0"]  # what was placed stays
    assert data["chartPages"] == [2]  # the four others are not lost


def test_a_rejected_crop_among_valid_ones_lists_the_chart_pages() -> None:
    def spoil_the_third(crops: list[dict]) -> list[dict]:
        return [{**c, "png": b"not PNG"} if c["chart"] == 2 else c for c in crops]

    data = _preview(_five_chart_workbook(), _EditedCrops(spoil_the_third))

    assert _ids(data) == ["chart_0", "chart_1", "chart_3", "chart_4"]
    assert data["chartPages"] == [2]


def test_every_crop_placed_lists_no_chart_pages() -> None:
    data = _preview(_five_chart_workbook(), _GalleryHost())

    assert _ids(data) == [f"chart_{n}" for n in range(5)]
    assert data["chartPages"] is None


def test_a_placeholder_left_without_its_picture_lists_the_chart_pages() -> None:
    from test_xlsx_gallery import _chart_workbook
    from test_xlsx_placeholders import _adding, _chartex_frame, _two_cell, _with_drawing

    data = _with_drawing(
        _chart_workbook(), _adding(_two_cell(_chartex_frame(), (0, 40), (4, 46)))
    )
    host = _EditedCrops(lambda crops: [c for c in crops if c["chart"] == 0])  # the chart only

    preview = _preview(data, host)

    assert _ids(preview) == ["chart_0"]
    assert preview["chartPages"] == [2]


# --- a crop is placed only when it is a picture -----------------------------------------


def test_a_png_that_is_only_named_so_is_not_placed() -> None:
    host = _GalleryHost(png=b"not PNG")

    data = _preview(_five_chart_workbook(), host)

    assert _ids(data) == []
    assert data["chartPages"] == [2]


def test_real_pngs_are_placed_as_data_uris() -> None:
    data = _preview(_five_chart_workbook(), _GalleryHost(png=_png(3, 5)))

    src = data["workbook"]["sheet"][0]["images"][0]["src"]
    assert src.startswith("data:image/png;base64,")


# --- cropping a gallery page --------------------------------------------------------------

_OUTLINE = "0 0 0 RG 1 w 100 300 200 150 re S"
_WHOLE = f"BT /F1 6 Tf 20 800 Td (lcx0123456789abchart1) Tj ET {_OUTLINE}"
# The second fragment starts where the first ends (Helvetica 6pt), so the page text
# still reads as one marker, the way LibreOffice splits a long cell text.
_SPLIT = (
    "BT /F1 6 Tf 20 800 Td (lcx0123456789ab) Tj ET "
    f"BT /F1 6 Tf 67.12 800 Td (chart1) Tj ET {_OUTLINE}"
)


def _crop(content: str) -> dict:
    (crop,) = PdfSourceConverter().chart_images(_vector_pdf([content]), path="gallery.pdf")
    return crop


def test_a_marker_split_over_text_objects_stays_out_of_the_crop() -> None:
    whole, split = _crop(_WHOLE), _crop(_SPLIT)

    assert (split["chart"], split["salt"]) == (1, "0123456789ab")
    assert all(abs(a - b) <= 1 for a, b in zip(whole["box"], split["box"], strict=True))
    assert abs(split["box"][2] - 202) <= 2 and abs(split["box"][3] - 152) <= 2


def test_text_beside_the_chart_that_only_looks_like_a_marker_fragment_stays_in() -> None:
    # "1" and "chart" are pieces of the marker string, but sit away from it: axis labels
    label = "BT /F1 10 Tf 40 250 Td (1) Tj ET BT /F1 10 Tf 60 250 Td (chart) Tj ET"

    crop = _crop(f"{_SPLIT} {label}")

    x, y, width, height = crop["box"]
    assert x <= 45 and y + height >= 590  # the labels widen the box; the marker does not
    assert y > 100  # the marker, at the page top, is still left out


def _rotated(content: str, degrees: int) -> bytes:
    document = pdfium.PdfDocument(_vector_pdf([content]))
    document[0].set_rotation(degrees)
    out = io.BytesIO()
    document.save(out)
    document.close()
    return out.getvalue()


def test_a_rotated_gallery_page_is_not_cropped() -> None:
    assert len(PdfSourceConverter().chart_images(_rotated(_WHOLE, 0), path="g.pdf")) == 1
    for degrees in (90, 180, 270):
        crops = PdfSourceConverter().chart_images(_rotated(_WHOLE, degrees), path="g.pdf")
        assert crops == [], degrees
