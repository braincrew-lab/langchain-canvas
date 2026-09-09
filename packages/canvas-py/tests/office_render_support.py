"""Measurement support for ``test_office_render_gate.py`` (test-only).

Runs the recorded Docker image's ``soffice`` (network disabled, production
arguments) and reads the resulting PDF with pdfium (``pypdfium2.raw``):
every glyph's tight box, font name and size, every drawn path's bounds,
segments, matrix and stroke width, and the comparison of those against the
frames the exported file declares.
Ported from ``docwriter-render-evidence/local-render/scripts/04_measure.py``
(verified on pdfium 149; pypdfium2 5.13.0 / pdfium 153 keeps the same raw
API names — a rename fails loudly as AttributeError, never a skip).

Assertions here are ``pytest.fail`` only: a failed conversion, a missing
PDF or a shape-count mismatch is a gate failure, not an environment skip.
"""

from __future__ import annotations

import ctypes
import io
import json
import math
import subprocess
import unicodedata
from pathlib import Path
from typing import Any

import pytest

IMAGE_ID = "sha256:aa6d1c9ddcbead7da95dce92f3bee979b1d9c4731abfe7b149b8ef19cd96951b"
ENGINE_LABEL = "LibreOffice 25.2.3 (Docker arm64, image aa6d1c9ddcbe)"
EMU_PER_POINT = 12700.0
EMU_PER_TWIP = 635
TOL = 1.0  # pt


# --- docker / soffice ---------------------------------------------------------------


def _docker(args: list[str], *, timeout: float = 55.0) -> subprocess.CompletedProcess[str]:
    """Run the recorded image with the network disabled; never a shell."""
    return subprocess.run(
        ["docker", "run", "--rm", "--network", "none", *args],
        capture_output=True,
        text=True,
        timeout=timeout,
        check=False,
    )


def _docker_bash(script: str) -> str:
    done = _docker(["--entrypoint", "/bin/bash", IMAGE_ID, "-c", script])
    if done.returncode != 0:
        pytest.fail(f"docker run failed ({done.returncode}): {done.stderr[-800:]}")
    return done.stdout


def render_to_pdf(path: Path) -> Path:
    """Convert ``path`` (inside the render dir) to PDF with the production command."""
    work = path.parent.resolve()
    pdf = path.with_suffix(".pdf")
    if pdf.exists():
        pdf.unlink()
    done = _docker(
        [
            "-v",
            f"{work}:/work",
            "--entrypoint",
            "soffice",
            IMAGE_ID,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nodefault",
            f"-env:UserInstallation=file:///tmp/lo-profile-{path.stem}",
            "--convert-to",
            "pdf",
            "--outdir",
            "/work",
            f"/work/{path.name}",
        ]
    )
    (work / f"{path.stem}.soffice.log").write_text(
        f"exit={done.returncode}\n--- stdout\n{done.stdout}\n--- stderr\n{done.stderr}",
        encoding="utf-8",
    )
    if done.returncode != 0:
        pytest.fail(f"soffice conversion of {path.name} failed: {done.stderr[-800:]}")
    if not pdf.exists():
        pytest.fail(f"soffice reported success but wrote no {pdf.name}")
    return pdf


# --- pdfium measurement (ported from docwriter-render-evidence 04_measure.py) --------


def _is_visible(ch: str) -> bool:
    if ch.isspace():
        return False
    category = unicodedata.category(ch)
    return not (category.startswith("C") or category.startswith("Z"))


def _hangul(ch: str) -> bool:
    return "가" <= ch <= "힣" or "ᄀ" <= ch <= "ᇿ" or "㄰" <= ch <= "㆏"


def _page_glyphs(page: Any, page_h: float) -> list[dict[str, Any]]:
    import pypdfium2.raw as raw

    textpage = page.get_textpage()
    count = textpage.count_chars()
    out: list[dict[str, Any]] = []
    buf = ctypes.create_string_buffer(256)
    flags = ctypes.c_int()
    left, bottom, right, top = (ctypes.c_double() for _ in range(4))
    origin_x, origin_y = ctypes.c_double(), ctypes.c_double()
    for index in range(count):
        code = raw.FPDFText_GetUnicode(textpage, index)
        ch = chr(code) if code else "\x00"
        raw.FPDFText_GetCharBox(textpage, index, left, right, bottom, top)
        raw.FPDFText_GetFontInfo(textpage, index, buf, 256, flags)
        size = raw.FPDFText_GetFontSize(textpage, index)
        # The glyph origin is the baseline; the tight box's bottom is not (a
        # Hangul glyph descends ~1 pt below the baseline a Latin digit sits on).
        if not raw.FPDFText_GetCharOrigin(textpage, index, origin_x, origin_y):
            pytest.fail(f"FPDFText_GetCharOrigin failed on glyph {index}")
        out.append(
            {
                "i": index,
                "ch": ch,
                "left": left.value,
                "right": right.value,
                "top": page_h - top.value,
                "bottom": page_h - bottom.value,
                "origin_x": origin_x.value,
                "origin_y": page_h - origin_y.value,
                "font": buf.value.decode("utf-8", "replace"),
                "size": size,
                "visible": _is_visible(ch),
            }
        )
    return out


SEGMENT_MOVETO = 2  # pdfium FPDF_SEGMENT_MOVETO
SEGMENT_LINETO = 0  # pdfium FPDF_SEGMENT_LINETO


def _matrix_uniform_scale(matrix: list[float]) -> float | None:
    """``|a| == |d|`` with no shear/rotation → that scale; otherwise ``None``."""
    a, b, c, d, _, _ = matrix
    if abs(b) > 1e-6 or abs(c) > 1e-6 or abs(abs(a) - abs(d)) > 1e-6:
        return None
    return abs(a)


def _page_paths(page: Any, page_h: float) -> list[dict[str, Any]]:
    """Every drawn path: page-space bounds (recorded only), every segment's
    point transformed by the object's matrix into page space (top-left
    origin, pt), the matrix, the stroke width and the draw mode — the raw
    material of R6's centreline identity."""
    import pypdfium2.raw as raw

    paths = []
    for obj in page.get_objects():
        if obj.type != raw.FPDF_PAGEOBJ_PATH:
            continue
        left, bottom, right, top = obj.get_bounds()
        matrix = raw.FS_MATRIX()
        if not raw.FPDFPageObj_GetMatrix(obj, matrix):
            pytest.fail("FPDFPageObj_GetMatrix failed on a path object")
        m = [matrix.a, matrix.b, matrix.c, matrix.d, matrix.e, matrix.f]
        stroke = ctypes.c_float()
        if not raw.FPDFPageObj_GetStrokeWidth(obj, stroke):
            pytest.fail("FPDFPageObj_GetStrokeWidth failed on a path object")
        fill_mode, stroked = ctypes.c_int(), raw.FPDF_BOOL()
        if not raw.FPDFPath_GetDrawMode(obj, fill_mode, stroked):
            pytest.fail("FPDFPath_GetDrawMode failed on a path object")
        line_cap = raw.FPDFPageObj_GetLineCap(obj) if hasattr(raw, "FPDFPageObj_GetLineCap") else -1
        segments = []
        for index in range(raw.FPDFPath_CountSegments(obj)):
            segment = raw.FPDFPath_GetPathSegment(obj, index)
            x, y = ctypes.c_float(), ctypes.c_float()
            if not raw.FPDFPathSegment_GetPoint(segment, x, y):
                pytest.fail(f"FPDFPathSegment_GetPoint failed on segment {index}")
            px = m[0] * x.value + m[2] * y.value + m[4]
            py = m[1] * x.value + m[3] * y.value + m[5]
            segments.append(
                {
                    "type": raw.FPDFPathSegment_GetType(segment),
                    "close": bool(raw.FPDFPathSegment_GetClose(segment)),
                    "point_pt": [round(px, 3), round(page_h - py, 3)],
                }
            )
        scale = _matrix_uniform_scale(m)
        paths.append(
            {
                "bbox_pt": [
                    round(left, 2),
                    round(page_h - top, 2),
                    round(right - left, 2),
                    round(top - bottom, 2),
                ],
                "matrix": [round(v, 6) for v in m],
                "uniform_scale": scale,
                "stroke_width_raw": round(stroke.value, 4),
                "stroke_width_pt": round(stroke.value * scale, 4) if scale is not None else None,
                "stroked": bool(stroked.value),
                "fill_mode": fill_mode.value,
                "line_cap": line_cap if line_cap >= 0 else None,
                "segments": segments,
            }
        )
    return paths


LINECAP_BUTT = 0  # pdfium FPDF_LINECAP_BUTT


def stroke_ink_box_pt(path: dict[str, Any]) -> list[float] | None:
    """The ink a stroked path leaves (top-left origin, pt): its segment points
    widened by half its stroke — across the line always, along it unless the
    cap is butt (an unknown cap gets the wider, square-cap box). pdfium's
    ``bbox_pt`` pads by the FULL stroke on every side (a 0.5 pt border reports
    a 1.0 pt tall box), which is not ink — the same reason R6 reads segment
    points rather than bounds. ``None`` for a filled (unstroked) path or one
    whose matrix has no uniform scale."""
    if not path["stroked"] or path["stroke_width_pt"] is None:
        return None
    points = [segment["point_pt"] for segment in path["segments"]]
    if not points:
        return None
    half = path["stroke_width_pt"] / 2
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    along = 0.0 if path.get("line_cap") == LINECAP_BUTT else half
    dx, dy = (along, half) if max(ys) - min(ys) <= max(xs) - min(xs) else (half, along)
    return [
        round(min(xs) - dx, 3),
        round(min(ys) - dy, 3),
        round(max(xs) - min(xs) + 2 * dx, 3),
        round(max(ys) - min(ys) + 2 * dy, 3),
    ]


def _axis_line(path: dict[str, Any], tol: float) -> tuple[str, float, float, float] | None:
    """``("h", x0, x1, y)`` / ``("v", y0, y1, x)`` for a stroked path whose
    points all lie on one horizontal / vertical line (within ``tol``)."""
    if stroke_ink_box_pt(path) is None:
        return None
    xs = [segment["point_pt"][0] for segment in path["segments"]]
    ys = [segment["point_pt"][1] for segment in path["segments"]]
    if max(ys) - min(ys) <= tol < max(xs) - min(xs):
        return "h", min(xs), max(xs), sum(ys) / len(ys)
    if max(xs) - min(xs) <= tol < max(ys) - min(ys):
        return "v", min(ys), max(ys), sum(xs) / len(xs)
    return None


def _union_boxes(boxes: list[list[float]]) -> list[float] | None:
    if not boxes:
        return None
    x0 = min(b[0] for b in boxes)
    y0 = min(b[1] for b in boxes)
    x1 = max(b[0] + b[2] for b in boxes)
    y1 = max(b[1] + b[3] for b in boxes)
    return [round(x0, 3), round(y0, 3), round(x1 - x0, 3), round(y1 - y0, 3)]


def table_grid_paths(
    paths: list[dict[str, Any]], min_width_pt: float = 300.0, tol: float = TOL
) -> dict[str, Any]:
    """The rendered table grid as a LATTICE: the horizontal strokes at least
    ``min_width_pt`` wide whose two ends each meet a vertical stroke (within
    ``tol``), plus every vertical those horizontals cross. A wide horizontal
    with no vertical at its ends — a paragraph rule — goes to ``other_wide``
    however its left x rounds. Also returns the union of the grid's ink boxes
    and its centre-line box (the verticals' x, the horizontals' y)."""
    geometry = [(index, _axis_line(path, tol)) for index, path in enumerate(paths)]
    verticals = [(i, g) for i, g in geometry if g is not None and g[0] == "v"]
    wide = [
        (i, g)
        for i, g in geometry
        if g is not None and g[0] == "h" and g[2] - g[1] >= min_width_pt
    ]

    def meets(x: float, y: float, vertical: tuple[str, float, float, float]) -> bool:
        _, y0, y1, vx = vertical
        return abs(vx - x) <= tol and y0 - tol <= y <= y1 + tol

    grid_index: set[int] = set()
    other: list[dict[str, Any]] = []
    horizontals = []
    for i, (_, x0, x1, y) in wide:
        left_end = any(meets(x0, y, g) for _, g in verticals)
        right_end = any(meets(x1, y, g) for _, g in verticals)
        if left_end and right_end:
            grid_index.add(i)
            horizontals.append((x0, x1, y))
        else:
            other.append(paths[i])
    for j, (_, y0, y1, vx) in verticals:
        crossed = (
            x0 - tol <= vx <= x1 + tol and y0 - tol <= y <= y1 + tol for x0, x1, y in horizontals
        )
        if any(crossed):
            grid_index.add(j)
    grid = [paths[i] for i in sorted(grid_index)]
    ink = [box for box in (stroke_ink_box_pt(p) for p in grid) if box is not None]
    centreline = None
    if grid:
        vxs = [g[3] for j, g in verticals if j in grid_index]
        hys = [y for _, _, y in horizontals]
        centreline = [
            round(min(vxs), 3),
            round(min(hys), 3),
            round(max(vxs) - min(vxs), 3),
            round(max(hys) - min(hys), 3),
        ]
    return {
        "grid": grid,
        "other_wide": other,
        "grid_ink_pt": _union_boxes(ink),
        "centreline_pt": centreline,
    }


def marker_run_before(glyphs: list[dict[str, Any]], item: str, marker: str) -> list[dict[str, Any]]:
    """The list marker drawn for ``item``: the ``marker`` glyphs immediately
    before the item's glyphs in the page stream (LibreOffice emits the label,
    then the text). A ``marker`` anywhere else never stands in for it."""
    stream = "".join(g["ch"] for g in glyphs)
    target = "".join(c for c in item if _is_visible(c))
    label = "".join(c for c in marker if _is_visible(c))
    start = stream.find(target)
    if start == -1:
        pytest.fail(f"{item!r} was not found among the rendered glyphs")
    if start < len(label) or stream[start - len(label) : start] != label:
        pytest.fail(
            f"{marker!r} does not immediately precede {item!r} in the glyph stream: "
            f"{stream[max(0, start - 8) : start + len(target)]!r}"
        )
    return glyphs[start - len(label) : start]


def _is_two_point_stroke(path: dict[str, Any]) -> bool:
    segments = path["segments"]
    return (
        len(segments) == 2
        and segments[0]["type"] == SEGMENT_MOVETO
        and segments[1]["type"] == SEGMENT_LINETO
    )


def _segment_centre(path: dict[str, Any]) -> list[float] | None:
    points = [s["point_pt"] for s in path["segments"]]
    if not points:
        return None
    if _is_two_point_stroke(path):
        (x0, y0), (x1, y1) = points
        return [(x0 + x1) / 2, (y0 + y1) / 2]
    xs, ys = [p[0] for p in points], [p[1] for p in points]
    return [(min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2]


def connector_path_identity(
    paths: list[dict[str, Any]], mid_pt: list[float], tol: float = TOL
) -> dict[str, Any]:
    """R6 identity (plan §실제 렌더 게이트, task 1b): the ONE two-point stroke
    path whose transformed midpoint is within ``tol`` of the file connector's
    midpoint. A filled polygon or a non-uniform matrix at that spot is a
    harness decision, not a wider window: recorded and ``pytest.fail``."""
    near = []
    for index, path in enumerate(paths):
        centre = _segment_centre(path)
        if centre is None:
            continue
        if abs(centre[0] - mid_pt[0]) <= tol and abs(centre[1] - mid_pt[1]) <= tol:
            near.append((index, path))
    polygons = [(i, p) for i, p in near if not _is_two_point_stroke(p)]
    if polygons:
        pytest.fail(
            "the path at the connector midpoint is not a two-point stroke (harness "
            f"decision needed): {json.dumps([p for _, p in polygons])}"
        )
    matches = [(i, p) for i, p in near if _is_two_point_stroke(p)]
    if len(matches) != 1:
        pytest.fail(
            f"expected exactly one two-point stroke path with midpoint within {tol} pt "
            f"of {mid_pt}, found {len(matches)}: {json.dumps([p for _, p in matches])}"
        )
    index, path = matches[0]
    if path["uniform_scale"] is None:
        pytest.fail(
            f"the connector path's matrix is not a uniform scale (harness decision "
            f"needed): matrix={path['matrix']} segments={json.dumps(path['segments'])}"
        )
    (x0, y0), (x1, y1) = (s["point_pt"] for s in path["segments"])
    return {
        "path_index": index,
        "points_pt": [[x0, y0], [x1, y1]],
        "midpoint_pt": [(x0 + x1) / 2, (y0 + y1) / 2],
        "drop_pt": y1 - y0,
        "angle_deg": math.degrees(math.atan2(y1 - y0, x1 - x0)),
        "matrix": path["matrix"],
        "uniform_scale": path["uniform_scale"],
        "stroke_width_raw": path["stroke_width_raw"],
        "stroke_width_pt": path["stroke_width_pt"],
        "stroked": path["stroked"],
        "closed": path["segments"][1]["close"],
        "bbox_pt": path["bbox_pt"],
    }


def measure_pdf_pages(pdf: Path) -> list[dict[str, Any]]:
    """Every page: size, drawn path boxes, visible glyphs (top-left origin, pt)."""
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(pdf))
    pages = []
    for index in range(len(document)):
        page = document[index]
        width, height = page.get_size()
        glyphs = [g for g in _page_glyphs(page, height) if g["visible"]]
        paths = _page_paths(page, height)
        pages.append(
            {
                "page": index + 1,
                "size_pt": [width, height],
                "paths_pt": [p["bbox_pt"] for p in paths],
                "paths": paths,
                "glyphs": glyphs,
            }
        )
    return pages


def _match(stream: str, glyphs: list[dict[str, Any]], text: str, used: list[bool]):
    target = "".join(c for c in text if _is_visible(c))
    if not target:
        return None, [], 0
    start = stream.find(target)
    while start != -1 and any(used[start : start + len(target)]):
        start = stream.find(target, start + 1)
    if start == -1:
        best = 0
        for k in range(len(target), 0, -1):
            s = stream.find(target[:k])
            if s != -1 and not any(used[s : s + k]):
                best, start = k, s
                break
        if best == 0:
            return None, [], len(target)
        for j in range(start, start + best):
            used[j] = True
        return (start, best), glyphs[start : start + best], len(target) - best
    for j in range(start, start + len(target)):
        used[j] = True
    return (start, len(target)), glyphs[start : start + len(target)], 0


def _lines(glyphs: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    lines: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    for glyph in glyphs:
        if current:
            prev = current[-1]
            h = max(prev["bottom"] - prev["top"], 1.0)
            if (
                abs(glyph["bottom"] - prev["bottom"]) > 0.6 * h
                or glyph["left"] < prev["left"] - 2 * h
            ):
                lines.append(current)
                current = []
        current.append(glyph)
    if current:
        lines.append(current)
    return lines


def _break_kinds(
    text: str, span: tuple[int, int] | None, lines: list[list[dict[str, Any]]]
) -> list[dict[str, str]]:
    breaks = []
    visible_index = [k for k, c in enumerate(text) if _is_visible(c)]
    pos = 0
    for a, b in zip(lines, lines[1:], strict=False):
        pos += len(a)
        k = pos - 1
        last, first = a[-1]["ch"], b[0]["ch"]
        between = (
            text[visible_index[k] + 1 : visible_index[k + 1]] if k + 1 < len(visible_index) else ""
        )
        if any(c.isspace() for c in between):
            kind = "word"
        elif last in "-‐–—/":
            kind = "after-hyphen"
        elif _hangul(last) and _hangul(first):
            kind = "intra-word:hangul"
        elif last.isascii() and first.isascii():
            kind = "intra-word:latin"
        else:
            kind = "intra-word:other"
        breaks.append(
            {
                "after": "".join(g["ch"] for g in a)[-8:],
                "before": "".join(g["ch"] for g in b)[:8],
                "kind": kind,
            }
        )
    return breaks


def _union(glyphs: list[dict[str, Any]]) -> list[float]:
    left = min(g["left"] for g in glyphs)
    right = max(g["right"] for g in glyphs)
    top = min(g["top"] for g in glyphs)
    bottom = max(g["bottom"] for g in glyphs)
    return [left, top, right - left, bottom - top]


def _inside_page(union: list[float], size: list[float]) -> bool:
    x, y, w, h = union
    return x >= -TOL and y >= -TOL and x + w <= size[0] + TOL and y + h <= size[1] + TOL


def _bullet_char(paragraph: Any) -> str:
    from pptx.oxml.ns import qn

    bullet = paragraph._p.find(".//" + qn("a:buChar"))
    return (bullet.get("char") or "") if bullet is not None else ""


def _bullet_at(
    glyphs: list[dict[str, Any]], used: list[bool], span: tuple[int, int] | None, marker: str
) -> tuple[str, list[dict[str, Any]]]:
    """The glyphs immediately before a matched body run, accepted as the
    paragraph's bullet only when they ARE the ``a:buChar`` text and unused.
    Returns ``(what was found there, the accepted glyphs or [])``."""
    if span is None or not marker:
        return "", []
    start = span[0]
    begin = start - len(marker)
    if begin < 0:
        return "", []
    candidate = glyphs[begin:start]
    found = "".join(g["ch"] for g in candidate)
    if found != marker or any(used[begin:start]):
        return found, []
    for j in range(begin, start):
        used[j] = True
    return found, candidate


def _shape_record(
    shape: Any, element: dict[str, Any], page: dict[str, Any], used: list[bool], stream: str
) -> dict[str, Any]:
    from pptx.oxml.ns import qn

    frame = [
        shape.left / EMU_PER_POINT,
        shape.top / EMU_PER_POINT,
        shape.width / EMU_PER_POINT,
        shape.height / EMU_PER_POINT,
    ]
    record: dict[str, Any] = {"id": element["id"], "type": element["type"], "frame_pt": frame}
    if not (shape.has_text_frame and shape.text_frame.text):
        return record
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    autofit = next(
        (
            c.tag.split("}")[1]
            for c in body
            if c.tag.split("}")[1] in ("noAutofit", "spAutoFit", "normAutofit")
        ),
        "inherited",
    )
    autofit_el = body.find(qn(f"a:{autofit}")) if autofit != "inherited" else None
    paragraphs = shape.text_frame.paragraphs
    bullets = [_bullet_char(p) for p in paragraphs]
    text = "\n".join(p.text for p in paragraphs)
    # A paragraph bullet (a:buChar) is a list marker the renderer draws in its
    # OWN marker face (LibreOffice: OpenSymbol) — a real glyph on the page that
    # paragraph.text does not carry, and not a body-font substitution. It is a
    # bullet glyph only when it sits immediately before the paragraph's first
    # body glyph AND is that character (positive match); anything else at that
    # position is a missing glyph. Every body glyph stays in the body set, so
    # a body character drawn by another face surfaces in ``pdf_fonts``.
    text_glyphs: list[dict[str, Any]] = []
    bullet_glyphs: list[dict[str, Any]] = []
    mismatches: list[dict[str, str]] = []
    missing = 0
    for bullet, paragraph in zip(bullets, paragraphs, strict=True):
        span, run, run_missing = _match(stream, page["glyphs"], paragraph.text, used)
        missing += run_missing
        text_glyphs.extend(run)
        if not bullet:
            continue
        marker = "".join(c for c in bullet if _is_visible(c))
        found, hit = _bullet_at(page["glyphs"], used, span, marker)
        if hit:
            bullet_glyphs.extend(hit)
        else:
            missing += len(marker)
            mismatches.append({"expected": marker, "found": found})
    n_visible = sum(1 for c in text if _is_visible(c)) + sum(
        1 for b in bullets for c in b if _is_visible(c)
    )
    glyphs = text_glyphs + bullet_glyphs
    record.update(
        {
            "text": text,
            "bullets": [b for b in bullets if b],
            "autofit": autofit,
            "fontScale": autofit_el.get("fontScale") if autofit_el is not None else None,
            "declared_size_pt": sorted(
                {r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size}
            ),
            "n_visible_chars": n_visible,
            "matched_glyphs": len(glyphs),
            "missing_glyphs": missing,
            "coverage": (len(glyphs) / n_visible) if n_visible else None,
            "bullet_glyph_count": len(bullet_glyphs),
            "bullet_mismatches": mismatches,
            "bullet_fonts": sorted({g["font"] for g in bullet_glyphs}),
        }
    )
    if not glyphs:
        record.update(
            {
                "contained_in_declared_frame": False,
                "lines": [],
                "line_count": 0,
                "breaks": [],
                "pdf_fonts": [],
            }
        )
        return record
    fx, fy, fw, fh = frame
    union = _union(glyphs)
    left, top, width, height = union
    right, bottom = left + width, top + height
    # Lines cluster on the run glyphs only: a list bullet's tight box sits
    # above the baseline, which would split it off as a line of its own.
    lines = _lines(text_glyphs)
    breaks = _break_kinds(text, None, lines)
    record.update(
        {
            "glyph_union_pt": union,
            "overflow_pt": {
                "left": max(0.0, fx - left),
                "right": max(0.0, right - (fx + fw)),
                "top": max(0.0, fy - top),
                "bottom": max(0.0, bottom - (fy + fh)),
            },
            "h_contained_1pt": left >= fx - TOL and right <= fx + fw + TOL,
            "v_contained_1pt": top >= fy - TOL and bottom <= fy + fh + TOL,
            "glyph_union_inside_page": _inside_page(union, page["size_pt"]),
            "text_grown_height_pt": bottom - fy,
            "pdf_fonts": sorted({g["font"] for g in text_glyphs}),
            "pdf_font_sizes": sorted({round(g["size"], 2) for g in text_glyphs}),
            "lines": ["".join(g["ch"] for g in line) for line in lines],
            "line_count": len(lines),
            "breaks": breaks,
            "break_kinds": {
                k: sum(1 for b in breaks if b["kind"] == k)
                for k in sorted({b["kind"] for b in breaks})
            },
        }
    )
    record["contained_in_declared_frame"] = (
        record["h_contained_1pt"] and record["v_contained_1pt"] and missing == 0
    )
    return record


def measure_deck(pdf: Path, pptx_bytes: bytes, elements: list[dict[str, Any]]) -> dict[str, Any]:
    """Glyph/frame comparison of the single-slide parity deck."""
    from pptx import Presentation

    deck = Presentation(io.BytesIO(pptx_bytes))
    (slide,) = deck.slides
    shapes = list(slide.shapes)
    if len(shapes) != len(elements):
        pytest.fail(f"exporter wrote {len(shapes)} shapes for {len(elements)} elements")
    pages = measure_pdf_pages(pdf)
    if len(pages) != 1:
        pytest.fail(f"expected one PDF page, got {len(pages)}")
    page = pages[0]
    stream = "".join(g["ch"] for g in page["glyphs"])
    used = [False] * len(page["glyphs"])
    records = [
        _shape_record(shape, element, page, used, stream)
        for shape, element in zip(shapes, elements, strict=True)
    ]
    connector = next(s for s, e in zip(shapes, elements, strict=True) if e["id"] == "line")
    begin = [connector.begin_x / EMU_PER_POINT, connector.begin_y / EMU_PER_POINT]
    end = [connector.end_x / EMU_PER_POINT, connector.end_y / EMU_PER_POINT]
    return {
        "engine": ENGINE_LABEL,
        "pdf": pdf.name,
        "pdf_page_size_pt": page["size_pt"],
        "slide_size_pt": [deck.slide_width / EMU_PER_POINT, deck.slide_height / EMU_PER_POINT],
        "pdf_fonts": sorted({g["font"] for g in page["glyphs"]}),
        "n_visible_glyphs": len(page["glyphs"]),
        "unmatched_visible_glyphs": sum(1 for u in used if not u),
        "rendered_path_boxes_pt": page["paths_pt"],
        "rendered_paths": page["paths"],
        "connector_pt": {
            "begin": begin,
            "end": end,
            "mid": [(begin[0] + end[0]) / 2, (begin[1] + end[1]) / 2],
        },
        "line_width_pt": connector.line.width / EMU_PER_POINT,
        "shapes": {r["id"]: r for r in records},
    }


def measure_document(pdf: Path, docx_bytes: bytes) -> dict[str, Any]:
    """Page geometry, section margins and glyph union of a Word file's render."""
    from docx import Document

    document = Document(io.BytesIO(docx_bytes))
    section = document.sections[0]
    margins_twips = {
        side: int(getattr(section, f"{side}_margin") / EMU_PER_TWIP)
        for side in ("left", "right", "top", "bottom")
    }
    page_w_pt = section.page_width / EMU_PER_POINT
    page_h_pt = section.page_height / EMU_PER_POINT
    column = [section.left_margin / EMU_PER_POINT, page_w_pt - section.right_margin / EMU_PER_POINT]
    pages = measure_pdf_pages(pdf)
    all_glyphs = [g for p in pages for g in p["glyphs"]]
    return {
        "engine": ENGINE_LABEL,
        "pdf": pdf.name,
        "file_page_size_pt": [page_w_pt, page_h_pt],
        "file_margins_twips": margins_twips,
        "file_text_column_pt": column,
        "pdf_pages": [
            {
                "page": p["page"],
                "size_pt": p["size_pt"],
                "paths_pt": p["paths_pt"],
                "paths": p["paths"],
                "glyph_union_pt": _union(p["glyphs"]) if p["glyphs"] else None,
                "fonts": sorted({g["font"] for g in p["glyphs"]}),
            }
            for p in pages
        ],
        "glyph_union_all_pages_pt": _union(all_glyphs) if all_glyphs else None,
        "glyphs": all_glyphs,
        "tables_in_file": len(document.tables),
    }


def _glyph_run(glyphs: list[dict[str, Any]], needle: str) -> list[dict[str, Any]]:
    stream = "".join(g["ch"] for g in glyphs)
    target = "".join(c for c in needle if _is_visible(c))
    start = stream.find(target)
    if start == -1:
        pytest.fail(f"{needle!r} was not found among the rendered glyphs")
    return glyphs[start : start + len(target)]
