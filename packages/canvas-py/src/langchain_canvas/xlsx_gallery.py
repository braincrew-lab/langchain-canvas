"""The chart gallery: a print copy of a workbook with one marked page per chart.

The read-only grid shows each chart as a picture at its anchor cell, and a host
renders the pictures from this copy. Each page carries a small text
``lcx<salt>chart<k>``, so the crop of that page says which chart it is and no page
map is predicted. :func:`xlsx_chart_gallery` builds the copy;
``PdfSourceConverter.chart_images`` crops it; ``replay`` pairs the crops with the
chart boxes ``xlsx_import.xlsx_sheets_and_charts`` reports.
"""

from __future__ import annotations

import hashlib
import io
import re
from datetime import datetime
from typing import Any

from .xlsx_import import _EMU_PER_PX

#: Charts one workbook's gallery copy may hold; a workbook with more is not galleried.
MAX_GALLERY_CHARTS = 40
_MIN_GALLERY_PX = (120, 80)
_MAX_GALLERY_PX = (1600, 1200)
_CORE_DATE = re.compile(rb"(<dcterms:(created|modified)\b[^>]*>)[^<]*(</dcterms:\2>)")
_FIXED_CORE_DATE = b"2000-01-01T00:00:00Z"

#: What the grid shows in place of a chart the preview cannot draw (both languages, one line).
UNSUPPORTED_CHART_TEXT = (
    "미리보기에서 지원하지 않는 차트입니다 / Chart type not supported in preview"
)
_MAX_PLACEHOLDER_COLUMNS = 25
_MAX_PLACEHOLDER_ROWS = 60


def _deterministic_save(workbook: Any) -> bytes:
    """The workbook's bytes, the same for the same content whenever it is saved.

    openpyxl stamps ``modified`` with the clock on every save and zip entries with
    the save time, so two saves of one workbook differ. The document dates and the
    zip timestamps are fixed here so a host's content-hash render cache holds.
    """
    import zipfile

    fixed = datetime(2000, 1, 1)
    workbook.properties.created = fixed
    workbook.properties.modified = fixed
    saved = io.BytesIO()
    workbook.save(saved)
    source = zipfile.ZipFile(io.BytesIO(saved.getvalue()))
    out = io.BytesIO()
    with source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            stamped = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            stamped.compress_type = zipfile.ZIP_DEFLATED
            content = source.read(info)
            if info.filename == "docProps/core.xml":
                # save() overwrites ``modified`` with the clock after the line above
                content = _CORE_DATE.sub(rb"\g<1>" + _FIXED_CORE_DATE + rb"\g<3>", content)
            target.writestr(stamped, content)
    return out.getvalue()


def gallery_salt(data: bytes) -> str:
    """The stamp of a workbook's gallery: the first 12 hex digits of its sha256.

    Every gallery page carries it, so a crop is only paired with the workbook whose
    gallery it came from.
    """
    return hashlib.sha256(data).hexdigest()[:12]


def _gallery_size(box: dict[str, Any]) -> tuple[int, int]:
    """The px size a chart is drawn at in the gallery: its own, kept within bounds.

    Grown to at least 120x80 (a chart folded to a sliver still prints a readable page),
    shrunk into 1600x1200, both with the aspect kept; then each side is raised to its
    minimum on its own, which bends the aspect only for degenerate shapes.
    """
    width, height = max(float(box["width"]), 1.0), max(float(box["height"]), 1.0)
    grow = max(1.0, _MIN_GALLERY_PX[0] / width, _MIN_GALLERY_PX[1] / height)
    width, height = width * grow, height * grow
    shrink = min(1.0, _MAX_GALLERY_PX[0] / width, _MAX_GALLERY_PX[1] / height)
    return (
        max(_MIN_GALLERY_PX[0], round(width * shrink)),
        max(_MIN_GALLERY_PX[1], round(height * shrink)),
    )


def _gallery_chart(workbook: Any, entry: dict[str, Any]) -> Any:
    """The chart object an entry points at: ``(sheet, slot)`` in the loaded workbook."""
    sheet, slot = entry.get("sheet"), entry.get("slot")
    sheets = workbook.worksheets
    if not (isinstance(sheet, int) and 0 <= sheet < len(sheets)):
        raise ValueError(f"the workbook has no sheet {sheet}")
    charts = sheets[sheet]._charts
    if not (isinstance(slot, int) and 0 <= slot < len(charts)):
        raise ValueError(f"sheet {sheet} has no chart in slot {slot}")
    return charts[slot]


def _new_page(workbook: Any, marker: str) -> Any:
    """A fit-to-one-page sheet named ``marker`` whose A1 says so."""
    from openpyxl.styles import Font  # type: ignore[import-untyped]
    from openpyxl.worksheet.properties import (  # type: ignore[import-untyped]
        PageSetupProperties,
    )

    page = workbook.create_sheet(marker)
    page["A1"] = marker
    page["A1"].font = Font(size=4, color="FFFFFF")  # on the page, but not for the eye
    page.column_dimensions["A"].width = 2
    page.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    page.page_setup.fitToWidth = 1
    page.page_setup.fitToHeight = 1
    return page


def _add_gallery_page(workbook: Any, marker: str, chart: Any, size: tuple[int, int]) -> Any:
    """A sheet named and marked ``marker`` that prints ``chart`` alone on one page."""
    from openpyxl.drawing.spreadsheet_drawing import (  # type: ignore[import-untyped]
        AnchorMarker,
        OneCellAnchor,
    )
    from openpyxl.drawing.xdr import XDRPositiveSize2D  # type: ignore[import-untyped]

    page = _new_page(workbook, marker)
    width, height = size
    page.add_chart(
        chart,
        OneCellAnchor(
            _from=AnchorMarker(col=1, row=1),
            ext=XDRPositiveSize2D(width * _EMU_PER_PX, height * _EMU_PER_PX),
        ),
    )
    return page


def _add_placeholder_page(workbook: Any, marker: str, size: tuple[int, int]) -> Any:
    """A marked sheet whose B2 holds the notice in a bordered block about ``size`` px.

    The block spans default-width columns (64px) and default-height rows (20px), at most
    25 by 60, so a chart that covers the whole sheet still gets a bounded page.
    """
    from openpyxl.styles import Alignment, Border, Side  # type: ignore[import-untyped]

    page = _new_page(workbook, marker)
    columns = min(_MAX_PLACEHOLDER_COLUMNS, max(1, round(size[0] / 64)))
    rows = min(_MAX_PLACEHOLDER_ROWS, max(1, round(size[1] / 20)))
    page.merge_cells(start_row=2, start_column=2, end_row=1 + rows, end_column=1 + columns)
    page["B2"] = UNSUPPORTED_CHART_TEXT
    page["B2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin")
    for row in page.iter_rows(min_row=2, max_row=1 + rows, min_col=2, max_col=1 + columns):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return page


def xlsx_chart_gallery(data: bytes, entries: list[dict[str, Any]]) -> bytes:
    """A print copy with one page per chart, each page marked ``lcx<salt>chart<k>``.

    ``entries`` are the chart boxes of the visible sheets, in the order the host
    pairs crops with them (see ``xlsx_sheets_and_charts``); ``k`` is the position of
    a chart in that list, and the page's marker says so. The host renders the copy to
    PDF and crops each marked page (``PdfSourceConverter.chart_images``), so which chart
    a picture is comes off the page itself and nothing about the page layout is guessed.

    Every original sheet stays in the copy, hidden, because the charts still read their
    data from them; each chart moves alone onto its own fit-to-one-page sheet at its
    original size. The bytes are the same for the same input (see ``_deterministic_save``).
    An ``"unsupported"`` entry (a chart the preview cannot draw) gets a page with a bordered
    notice (:data:`UNSUPPORTED_CHART_TEXT`) in place of a chart. Raises ``ValueError`` when
    ``entries`` is empty, holds more than :data:`MAX_GALLERY_CHARTS`, has nothing to draw,
    or names a chart the workbook cannot supply.
    """
    from openpyxl import load_workbook  # type: ignore[import-untyped]

    from .converters import ensure_archive_within_limits

    if not 1 <= len(entries) <= MAX_GALLERY_CHARTS:
        raise ValueError(f"a gallery holds 1 to {MAX_GALLERY_CHARTS} charts, got {len(entries)}")
    ensure_archive_within_limits(data, path="gallery.xlsx")
    workbook = load_workbook(io.BytesIO(data))

    # Resolve every chart before any is detached: moving one must not shift the slots
    # of the others.
    charts = {
        number: _gallery_chart(workbook, entry)
        for number, entry in enumerate(entries)
        if entry.get("kind") == "chart"
    }
    if not any(entry.get("kind") in ("chart", "unsupported") for entry in entries):
        raise ValueError("the entries hold nothing to put in a gallery")
    moved = {id(chart) for chart in charts.values()}
    for sheet in workbook.worksheets:
        sheet._charts = [chart for chart in sheet._charts if id(chart) not in moved]
        sheet.sheet_state = "hidden"
        sheet.print_area = None
        sheet._print_rows = None
        sheet._print_cols = None
    for chartsheet in workbook.chartsheets:
        chartsheet.sheet_state = "hidden"

    salt = gallery_salt(data)
    pages = []
    for number, entry in enumerate(entries):
        marker = f"lcx{salt}chart{number}"
        if number in charts:
            pages.append(_add_gallery_page(workbook, marker, charts[number], _gallery_size(entry)))
        elif entry.get("kind") == "unsupported":
            pages.append(_add_placeholder_page(workbook, marker, _gallery_size(entry)))
    workbook.active = pages[0]  # the sheet object: an index would count the hidden originals
    return _deterministic_save(workbook)
