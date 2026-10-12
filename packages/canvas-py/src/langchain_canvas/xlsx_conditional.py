"""Conditional formats of a read-only sheet, evaluated against the values the file holds.

A read-only grid never changes its values, so each rule is decided once here
and written into the cells it matches: font colour, bold, italic, strike and
underline, a fill (``bg``), a colour-scale colour, a data bar (``lcx.bar``) or
an icon (``lcx.icon``). Fortune's own rule engine is not used: it knows four
comparisons (no ``>=``, ``<=`` or ``<>``), no priority or stop-if-true, and
its data bars never find a maximum.

Excel's precedence is kept: rules run in priority order, a property a
higher-priority rule set is not overwritten, ``stopIfTrue`` ends a cell's
rules, and a colour from the number format beats a rule's font colour. A rule
this module cannot decide (an expression outside the small evaluator, an
oversized range) is listed on the sheet under ``lcxSkipped`` instead of being
dropped silently.
"""

from __future__ import annotations

import math
import re
from collections.abc import Callable
from datetime import date, datetime, timedelta
from statistics import mean, pstdev
from typing import Any

from openpyxl.worksheet.cell_range import CellRange  # type: ignore[import-untyped]

from .xlsx_import import _to_hex

#: Cells one rule may cover before it is skipped (and listed) instead of evaluated.
MAX_RULE_CELLS = 100_000

_REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_TOKEN = re.compile(
    r"\s*(?:"
    r"(?P<num>\d+(?:\.\d+)?(?:[eE][-+]?\d+)?)"
    r'|(?P<str>"(?:[^"]|"")*")'
    r"|(?P<ref>\$?[A-Z]{1,3}\$?\d+)"
    r"|(?P<name>[A-Z][A-Z0-9_.]*)(?=\s*\()"
    r"|(?P<bool>TRUE|FALSE)\b"
    r"|(?P<op><=|>=|<>|[-+*/^&=<>(),])"
    r")",
    re.I,
)


class Unsupported(Exception):
    """An expression the evaluator does not handle."""


def _column_number(letters: str) -> int:
    number = 0
    for char in letters.upper():
        number = number * 26 + ord(char) - 64
    return number


# --- a small expression evaluator ---------------------------------------------


def _evaluate(text: str, read: Callable[[int, int], Any], shift: tuple[int, int]) -> Any:
    """Evaluate an Excel formula written for the rule's top-left cell at a
    cell ``shift`` rows/columns away: relative references move, ``$`` ones stay."""
    tokens: list[tuple[str, str]] = []
    position = 0
    body = text[1:] if text.startswith("=") else text
    while position < len(body):
        found = _TOKEN.match(body, position)
        if not found or found.end() == position:
            if body[position:].strip() == "":
                break
            raise Unsupported(body[position:])
        kind = found.lastgroup or ""
        tokens.append((kind, found.group(kind)))
        position = found.end()

    index = 0

    def peek() -> tuple[str, str] | None:
        return tokens[index] if index < len(tokens) else None

    def take() -> tuple[str, str]:
        nonlocal index
        if index >= len(tokens):
            raise Unsupported("unexpected end")
        index += 1
        return tokens[index - 1]

    def primary() -> Any:
        kind, value = take()
        if kind == "num":
            return float(value)
        if kind == "str":
            return value[1:-1].replace('""', '"')
        if kind == "bool":
            return value.upper() == "TRUE"
        if kind == "ref":
            found = _REF.fullmatch(value.upper())
            assert found
            row = int(found.group(4)) + (0 if found.group(3) else shift[0])
            column = _column_number(found.group(2)) + (0 if found.group(1) else shift[1])
            return read(row, column)
        if kind == "name":
            take()  # (
            args: list[Any] = []
            if peek() != ("op", ")"):
                args.append(comparison())
                while peek() == ("op", ","):
                    take()
                    args.append(comparison())
            take()  # )
            return _call(value.upper(), args)
        if value == "(":
            inner = comparison()
            take()
            return inner
        if value == "-":
            return -_number(primary())
        if value == "+":
            return _number(primary())
        raise Unsupported(value)

    def power() -> Any:
        left = primary()
        while peek() == ("op", "^"):
            take()
            left = _number(left) ** _number(primary())
        return left

    def term() -> Any:
        left = power()
        while peek() in (("op", "*"), ("op", "/")):
            op = take()[1]
            right = _number(power())
            left = _number(left) * right if op == "*" else _number(left) / right
        return left

    def additive() -> Any:
        left = term()
        while peek() in (("op", "+"), ("op", "-")):
            op = take()[1]
            right = _number(term())
            left = _number(left) + right if op == "+" else _number(left) - right
        return left

    def concat() -> Any:
        left = additive()
        while peek() == ("op", "&"):
            take()
            left = _text(left) + _text(additive())
        return left

    def comparison() -> Any:
        left = concat()
        while peek() is not None and peek()[1] in ("=", "<>", "<", ">", "<=", ">="):  # type: ignore[index]
            op = take()[1]
            left = _compare(left, op, concat())
        return left

    result = comparison()
    if index != len(tokens):
        raise Unsupported("trailing tokens")
    return result


def _number(value: Any) -> float:
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise Unsupported(f"not a number: {value!r}") from exc


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _compare(left: Any, op: str, right: Any) -> bool:
    if isinstance(left, str) or isinstance(right, str):
        a: Any = _text(left).lower()
        b: Any = _text(right).lower()
    else:
        a, b = _number(left), _number(right)
    return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b}[op]


def _call(name: str, args: list[Any]) -> Any:
    if name == "AND":
        return all(_truthy(arg) for arg in args)
    if name == "OR":
        return any(_truthy(arg) for arg in args)
    if name == "NOT" and len(args) == 1:
        return not _truthy(args[0])
    if name == "ISBLANK" and len(args) == 1:
        return args[0] is None or args[0] == ""
    if name == "ABS" and len(args) == 1:
        return abs(_number(args[0]))
    if name == "MOD" and len(args) == 2:
        return math.fmod(_number(args[0]), _number(args[1]))
    if name == "LEN" and len(args) == 1:
        return float(len(_text(args[0])))
    raise Unsupported(name)


def _truthy(value: Any) -> bool:
    if isinstance(value, str):
        return value.upper() == "TRUE"
    return bool(_number(value))


# --- rule decisions -----------------------------------------------------------


def _numeric(value: Any) -> float | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return float(value)


def _cfvo_value(cfvo: Any, numbers: list[float]) -> float | None:
    kind = getattr(cfvo, "type", None)
    raw: Any = getattr(cfvo, "val", None)
    if not numbers:
        return None
    if kind == "min":
        return min(numbers)
    if kind == "max":
        return max(numbers)
    if raw is None and kind in ("num", "percent", "percentile"):
        raise Unsupported(f"cfvo {kind} without a value")
    if kind == "num":
        return float(raw)
    if kind == "percent":
        low, high = min(numbers), max(numbers)
        return low + (high - low) * float(raw) / 100
    if kind == "percentile":
        ordered = sorted(numbers)
        rank = (len(ordered) - 1) * float(raw) / 100
        lower = math.floor(rank)
        upper = min(lower + 1, len(ordered) - 1)
        return ordered[lower] + (ordered[upper] - ordered[lower]) * (rank - lower)
    raise Unsupported(f"cfvo {kind}")


def _mix(a: str, b: str, ratio: float) -> str:
    x, y = int(a[1:], 16), int(b[1:], 16)
    channels = [
        round(((x >> s) & 255) + (((y >> s) & 255) - ((x >> s) & 255)) * ratio) for s in (16, 8, 0)
    ]
    return "#" + "".join(f"{c:02X}" for c in channels)


def _scale_colour(value: float, stops: list[tuple[float, str]]) -> str:
    if value <= stops[0][0]:
        return stops[0][1]
    for (low, low_colour), (high, high_colour) in zip(stops, stops[1:], strict=False):
        if value <= high:
            span = high - low
            return _mix(low_colour, high_colour, 0 if span == 0 else (value - low) / span)
    return stops[-1][1]


def _period_holds(kind: str, value: Any, today: date) -> bool:
    if isinstance(value, datetime):
        day = value.date()
    elif isinstance(value, date):
        day = value
    else:
        return False
    start_week = today - timedelta(days=(today.weekday() + 1) % 7)
    month = (today.year, today.month)
    previous = (today.year - 1, 12) if today.month == 1 else (today.year, today.month - 1)
    following = (today.year + 1, 1) if today.month == 12 else (today.year, today.month + 1)
    return {
        "today": day == today,
        "yesterday": day == today - timedelta(days=1),
        "tomorrow": day == today + timedelta(days=1),
        "last7Days": today - timedelta(days=6) <= day <= today,
        "thisWeek": start_week <= day < start_week + timedelta(days=7),
        "lastWeek": start_week - timedelta(days=7) <= day < start_week,
        "nextWeek": start_week + timedelta(days=7) <= day < start_week + timedelta(days=14),
        "thisMonth": (day.year, day.month) == month,
        "lastMonth": (day.year, day.month) == previous,
        "nextMonth": (day.year, day.month) == following,
    }.get(kind, False)


def _colour_or_none(color: Any) -> str | None:
    if color is None:
        return None
    if getattr(color, "type", None) == "rgb" and getattr(color, "rgb", "") in ("00000000", None):
        return None
    return _to_hex(color)


def _dxf_style(dxf: Any) -> dict[str, Any]:
    """A differential style as the Fortune cell keys it sets."""
    style: dict[str, Any] = {}
    if dxf is None:
        return style
    font = getattr(dxf, "font", None)
    if font is not None:
        colour = _colour_or_none(getattr(font, "color", None))
        if colour:
            style["fc"] = colour
        if font.b:
            style["bl"] = 1
        if font.i:
            style["it"] = 1
        if font.strike:
            style["cl"] = 1
        if font.u:
            style["un"] = 1
    fill = getattr(dxf, "fill", None)
    if fill is not None and getattr(fill, "patternType", "solid") in (None, "solid"):
        colour = _colour_or_none(getattr(fill, "bgColor", None)) or _colour_or_none(
            getattr(fill, "fgColor", None)
        )
        if colour:
            style["bg"] = colour
    return style


class _Sheet:
    """Values and Fortune cells of one sheet, by 1-based position."""

    def __init__(self, sheet: dict[str, Any], cached_sheet: Any) -> None:
        self.sheet = sheet
        self.cached = cached_sheet
        self.cells = {(e["r"] + 1, e["c"] + 1): e["v"] for e in sheet["celldata"]}

    def value(self, row: int, column: int) -> Any:
        if row < 1 or column < 1:
            return None
        return self.cached.cell(row=row, column=column).value

    def cell(self, row: int, column: int) -> dict[str, Any]:
        v = self.cells.get((row, column))
        if v is None:
            v = {"m": "", "v": ""}
            self.sheet["celldata"].append({"r": row - 1, "c": column - 1, "v": v})
            self.cells[(row, column)] = v
        return v


def _positions(sqref: Any) -> list[tuple[int, int]]:
    return [
        (row, column)
        for cell_range in sqref.ranges
        for row in range(cell_range.min_row, cell_range.max_row + 1)
        for column in range(cell_range.min_col, cell_range.max_col + 1)
    ]


def _matches(
    rule: Any, sheet: _Sheet, positions: list[tuple[int, int]], top_left: tuple[int, int]
) -> dict[tuple[int, int], dict[str, Any]]:
    """Cells a rule applies to, with the keys it sets on each."""
    kind = rule.type
    values = {at: sheet.value(*at) for at in positions}
    numbers = [n for n in (_numeric(v) for v in values.values()) if n is not None]
    style = _dxf_style(rule.dxf)
    hits: dict[tuple[int, int], dict[str, Any]] = {}

    def shift(at: tuple[int, int]) -> tuple[int, int]:
        return at[0] - top_left[0], at[1] - top_left[1]

    if kind == "cellIs":
        formulas = list(rule.formula or [])
        for at, value in values.items():
            if value is None or value == "":
                continue
            bounds = [_evaluate(f, sheet.value, shift(at)) for f in formulas]
            op = rule.operator
            if op in ("between", "notBetween"):
                low, high = sorted((_number(bounds[0]), _number(bounds[1])))
                inside = _numeric(value) is not None and low <= float(value) <= high
                ok = inside if op == "between" else not inside
            else:
                symbol = {
                    "lessThan": "<",
                    "lessThanOrEqual": "<=",
                    "greaterThan": ">",
                    "greaterThanOrEqual": ">=",
                    "equal": "=",
                    "notEqual": "<>",
                }[op]
                ok = _compare(value, symbol, bounds[0])
            if ok:
                hits[at] = style
    elif kind == "expression":
        for at in positions:
            if _truthy(_evaluate(rule.formula[0], sheet.value, shift(at))):
                hits[at] = style
    elif kind in ("containsText", "notContainsText", "beginsWith", "endsWith"):
        needle = (rule.text or "").lower()
        for at, value in values.items():
            text = _text(value).lower()
            ok = {
                "containsText": needle in text,
                "notContainsText": needle not in text,
                "beginsWith": text.startswith(needle),
                "endsWith": text.endswith(needle),
            }[kind]
            if ok:
                hits[at] = style
    elif kind in ("containsBlanks", "notContainsBlanks"):
        for at, value in values.items():
            blank = value is None or _text(value).strip() == ""
            if blank == (kind == "containsBlanks"):
                hits[at] = style
    elif kind in ("containsErrors", "notContainsErrors"):
        for at, value in values.items():
            error = (
                isinstance(value, str) and value.startswith("#") and value.endswith(("!", "?", "A"))
            )
            if error == (kind == "containsErrors"):
                hits[at] = style
    elif kind in ("duplicateValues", "uniqueValues"):
        counts: dict[str, int] = {}
        for value in values.values():
            if value is not None and value != "":
                counts[_text(value).lower()] = counts.get(_text(value).lower(), 0) + 1
        for at, value in values.items():
            if value is None or value == "":
                continue
            duplicate = counts[_text(value).lower()] > 1
            if duplicate == (kind == "duplicateValues"):
                hits[at] = style
    elif kind == "top10":
        ordered = sorted(numbers, reverse=not rule.bottom)
        count = int(rule.rank or 10)
        if rule.percent:
            count = max(1, math.floor(len(ordered) * count / 100))
        if ordered:
            cut = ordered[min(count, len(ordered)) - 1]
            for at, value in values.items():
                number = _numeric(value)
                if number is not None and (number <= cut if rule.bottom else number >= cut):
                    hits[at] = style
    elif kind == "aboveAverage":
        if numbers:
            average = mean(numbers)
            spread = pstdev(numbers) * (rule.stdDev or 0)
            above = rule.aboveAverage is not False
            for at, value in values.items():
                number = _numeric(value)
                if number is None:
                    continue
                if above:
                    ok = (
                        number >= average + spread
                        if rule.equalAverage
                        else number > average + spread
                    )
                else:
                    ok = (
                        number <= average - spread
                        if rule.equalAverage
                        else number < average - spread
                    )
                if ok:
                    hits[at] = style
    elif kind == "timePeriod":
        today = date.today()
        for at, value in values.items():
            if _period_holds(rule.timePeriod, value, today):
                hits[at] = style
    elif kind == "colorScale":
        scale = rule.colorScale
        stops = [
            (_cfvo_value(cfvo, numbers), _to_hex(colour) or "#FFFFFF")
            for cfvo, colour in zip(scale.cfvo, scale.color, strict=False)
        ]
        if numbers and all(point is not None for point, _ in stops):
            ordered_stops = sorted(((float(p), c) for p, c in stops), key=lambda s: s[0])  # type: ignore[arg-type]
            for at, value in values.items():
                number = _numeric(value)
                if number is not None:
                    hits[at] = {"bg": _scale_colour(number, ordered_stops)}
    elif kind == "dataBar":
        bar = rule.dataBar
        bar_low = _cfvo_value(bar.cfvo[0], numbers)
        bar_high = _cfvo_value(bar.cfvo[1], numbers)
        colour = _to_hex(bar.color) or "#638EC6"
        if bar_low is not None and bar_high is not None:
            span = (bar_high - bar_low) or 1
            for at, value in values.items():
                number = _numeric(value)
                if number is not None:
                    ratio = max(0.0, min(1.0, (number - bar_low) / span))
                    hits[at] = {"lcx": {"bar": {"ratio": round(ratio, 4), "color": colour}}}
    elif kind == "iconSet":
        icons = rule.iconSet
        name = icons.iconSet or "3TrafficLights1"
        bounds = [_cfvo_value(cfvo, numbers) for cfvo in icons.cfvo]
        if numbers and all(b is not None for b in bounds):
            for at, value in values.items():
                number = _numeric(value)
                if number is None:
                    continue
                bucket = 0
                for position, bound in enumerate(bounds[1:], start=1):
                    if number >= bound:  # type: ignore[operator]
                        bucket = position
                if icons.reverse:
                    bucket = len(bounds) - 1 - bucket
                hits[at] = {"lcx": {"icon": {"set": name, "index": bucket}}}
    else:
        raise Unsupported(kind)
    return hits


def bake_conditional_formats(
    sheet: dict[str, Any], worksheet: Any, cached_sheet: Any, format_coloured: set[tuple[int, int]]
) -> None:
    """Write a read-only sheet's conditional formats into its cells.

    ``format_coloured`` holds 0-based positions whose colour came from the
    number format; a rule's font colour does not override those.
    """
    rules: list[tuple[int, Any, Any]] = []
    for formatting in worksheet.conditional_formatting:
        for rule in formatting.rules:
            rules.append((rule.priority or 0, formatting.sqref, rule))
    if not rules:
        return
    rules.sort(key=lambda item: item[0])

    view = _Sheet(sheet, cached_sheet)
    claimed: dict[tuple[int, int], set[str]] = {}
    stopped: set[tuple[int, int]] = set()
    skipped: list[dict[str, str]] = []
    for _priority, sqref, rule in rules:
        positions = _positions(sqref)
        if len(positions) > MAX_RULE_CELLS:
            skipped.append({"type": rule.type, "range": str(sqref), "reason": "range too large"})
            continue
        first = CellRange(str(sqref).split()[0])
        try:
            hits = _matches(rule, view, positions, (first.min_row, first.min_col))
        except (Unsupported, ArithmeticError, IndexError, TypeError, ValueError) as exc:
            skipped.append({"type": rule.type, "range": str(sqref), "reason": str(exc)[:120]})
            continue
        for at, style in hits.items():
            if at in stopped:
                continue
            cell = view.cell(*at)
            taken = claimed.setdefault(at, set())
            for key, value in style.items():
                if key in taken:
                    continue
                if key == "fc" and (at[0] - 1, at[1] - 1) in format_coloured:
                    continue
                if key == "lcx":
                    cell["lcx"] = {**cell.get("lcx", {}), **value}
                    taken.update(f"lcx.{k}" for k in value)
                    continue
                cell[key] = value
                taken.add(key)
                if key == "bg":
                    # A rule's fill covers the cell's own patterned fill.
                    lcx = cell.get("lcx")
                    if isinstance(lcx, dict) and "fill" in lcx:
                        lcx = {k: v for k, v in lcx.items() if k != "fill"}
                        if lcx:
                            cell["lcx"] = lcx
                        else:
                            cell.pop("lcx")
            if rule.stopIfTrue:
                stopped.add(at)
    if skipped:
        sheet["lcxSkipped"] = skipped
